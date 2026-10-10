"""Part 6 tests: decision assistant, agent chat, debate, what-if, scenarios,
challenges, reports, history and the Decision Intelligence page.

Temporary database copies only. Requirements come from real G60 clusters
(rule-based); human assessment scores are SYNTHETIC TEST VALUES."""
from __future__ import annotations

import csv
import io
import json
import shutil
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from core.database import ensure_ai_schema, get_engine, session_scope
from core.decision_models import DecisionMessage, ManagementReportExport, PrioritizationScenario
from core.prioritization_models import AgentExecutionResult, PriorityAssessment, PrioritizationRun, RequirementPriorityScore
from services import agent_chat, agent_debate as deb, decision_assistant as da, decision_llm as dl
from services import agent_configuration as ac, dynamic_prioritization as dp, management_reports as mr
from services import priority_challenges as pc, prioritization_profiles as pp, requirement_review as rr
from services import scenario_comparison as sc, what_if_simulator as wis
from services.cross_source_mapping import build_cross_source_links
from services.decision_grounding import ref_exists
from services.insight_generator import generate_insights
from services.need_clustering import build_clusters, vfc_cluster_id
from services.need_taxonomy import sync_taxonomy
from services.requirement_evidence import build_evidence_package
from services.requirement_generator import generate_rule_based

ROOT = Path(__file__).resolve().parent.parent
USER = "test.pm"
NAMES = ["Touch screen, operation", "Electric Range", "Head-Up Display"]


@pytest.fixture(scope="module")
def base_db(built, tmp_path_factory):
    path = tmp_path_factory.mktemp("p6") / "base.db"
    shutil.copy2(built["db_path"], path)
    engine = get_engine(path)
    ensure_ai_schema(engine)
    with session_scope(engine) as s:
        sync_taxonomy(s)
        build_cross_source_links(s)
        build_clusters(s)
        generate_insights(s, min_unique_feedback=10)
    engine.dispose()
    return path


@pytest.fixture(scope="module")
def run_db(base_db, tmp_path_factory):
    """3 requirements (real clusters), 4 agents incl. a custom human agent, a
    run where Head-Up Display is INCOMPLETE (no human assessments)."""
    path = tmp_path_factory.mktemp("p6") / "run.db"
    shutil.copy2(base_db, path)
    engine = get_engine(path)
    with session_scope(engine) as s:
        reqs = {n: rr.create_requirement(s, generate_rule_based(build_evidence_package(s, cluster_id=vfc_cluster_id(n))),
                                         USER).requirement_id for n in NAMES}
        dem, _ = ac.create_agent(s, ac.STARTER_AGENTS[0].model_copy(), USER)
        cov, _ = ac.create_agent(s, ac.STARTER_AGENTS[1].model_copy(), USER)
        imp, _ = ac.create_agent(s, ac.STARTER_AGENTS[2].model_copy(), USER)
        tf, _ = ac.create_agent(s, ac.AgentSpec(name="Technology Feasibility Agent", objective="Custom feasibility view",
                                                execution_method="HUMAN_ASSESSMENT", scoring_method="HUMAN_SCORE",
                                                resources=["RES_HUMAN_ASSESSMENTS", "RES_REQUIREMENTS"]), USER)
        p = pp.create_profile(s, "P6 test", {dem.agent_id: 30, cov.agent_id: 20, imp.agent_id: 25, tf.agent_id: 25}, USER)
        pp.activate_profile(s, p.profile_id, USER)
        for n, (i, t) in {"Touch screen, operation": (80, 60), "Electric Range": (70, 50)}.items():  # SYNTHETIC
            dp.record_human_assessment(s, reqs[n], imp.agent_id, i, "Synthetic test rationale (impact)", "test.assessor")
            dp.record_human_assessment(s, reqs[n], tf.agent_id, t, "Synthetic test rationale (feasibility)", "test.engineer")
        run = dp.run_prioritization(s, p.profile_id, USER)
        meta = {"reqs": reqs, "run": run.run_id, "profile": p.profile_id,
                "agents": {"demand": dem.agent_id, "coverage": cov.agent_id, "impact": imp.agent_id, "feasibility": tf.agent_id}}
    engine.dispose()
    (path.parent / "meta.json").write_text(json.dumps(meta))
    return path


@pytest.fixture()
def env(run_db, tmp_path):
    path = tmp_path / "p6.db"
    shutil.copy2(run_db, path)
    meta = json.loads((run_db.parent / "meta.json").read_text())
    engine = get_engine(path)
    with session_scope(engine) as s:
        yield s, meta, path
    engine.dispose()


@pytest.fixture(autouse=True)
def no_llm(monkeypatch):
    for var in ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN", "PRODUCTMIND_ALLOW_EXTERNAL_LLM"):
        monkeypatch.delenv(var, raising=False)

    def forbid(*a, **k):
        raise AssertionError("No external LLM client may be constructed")

    monkeypatch.setattr("services.llm_client.AnthropicLLMClient.__init__", forbid)
    monkeypatch.setattr("anthropic.Anthropic.__init__", forbid, raising=False)
    dl.register_local_model(None)


def official_snapshot(s):
    return (sorted((r.requirement_id, r.final_score, r.rank, r.priority_status) for r in s.scalars(select(RequirementPriorityScore))),
            s.scalar(select(func.count()).select_from(AgentExecutionResult)), s.scalar(select(func.count()).select_from(PrioritizationRun)),
            sorted((a.assessment_id, a.score, a.superseded) for a in s.scalars(select(PriorityAssessment))),
            [(p.profile_id, p.version, p.status) for p in pp.list_profiles(s)])


# --- 1-3 assistant offline, no LLM, valid evidence ------------------------------------------------


@pytest.mark.parametrize("intent", ["explain_priority", "evidence", "missing", "challenge", "reliability"])
def test_offline_assistant_answers_without_api_key(env, intent):
    s, m, _ = env
    for rid in m["reqs"].values():
        resp = da.answer(s, rid, m["run"], intent=intent)
        assert resp.mode == "OFFLINE_EVIDENCE_BASED" and "not AI-generated" in resp.label
        assert resp.run_id == m["run"] and resp.items
        assert {i.kind for i in resp.items} <= {"FACT", "ASSESSMENT", "ASSUMPTION", "LIMITATION", "SIMULATION"}


def test_free_text_routing_and_clarification(env):
    s, m, _ = env
    rid = m["reqs"]["Touch screen, operation"]
    cases = {"Why is this requirement ranked first?": "explain_priority",
             "Why did Customer Demand Agent assign this score?": "agent_discussion",
             "Which customer comments support this requirement?": "evidence",
             "Why is this requirement incomplete?": "missing",
             "What happens if Customer Impact Agent receives 40% weight?": "scenario",
             "What if this customer problem is more serious than its ranking suggests?": "challenge",
             "How reliable is this recommendation?": "reliability"}
    for q, intent in cases.items():
        assert da.answer(s, rid, m["run"], question=q).intent == intent, q
    unclear = da.answer(s, rid, m["run"], question="tell me something nice")
    assert unclear.needs_clarification and unclear.suggestions and not [i for i in unclear.items if i.kind == "FACT"]


def test_every_referenced_evidence_id_is_valid(env):
    s, m, _ = env
    for rid in m["reqs"].values():
        for intent in ("evidence", "explain_priority"):
            resp = da.answer(s, rid, m["run"], intent=intent)
            assert all(ref_exists(s, r["table"], r["key"]) for r in resp.evidence_refs)
            for item in resp.items:
                for ref in item.refs:
                    table = "feedback" if ref.startswith("G60-") else None
                    if table:
                        assert ref_exists(s, table, ref)
    resp = da.answer(s, m["reqs"]["Touch screen, operation"], m["run"], intent="evidence")
    assert len([r for r in resp.evidence_refs if r["table"] == "feedback"]) == 70
    assert all(r["key"].startswith("G60-") for r in resp.evidence_refs if r["table"] == "feedback")


def test_answers_are_bound_to_one_run(env):
    s, m, _ = env
    with pytest.raises(ValueError):
        da.answer(s, m["reqs"]["Touch screen, operation"], "RUN-DOES-NOT-EXIST", intent="explain_priority")
    resp = da.answer(s, m["reqs"]["Touch screen, operation"], m["run"], intent="explain_priority")
    run = s.get(PrioritizationRun, m["run"])
    assert resp.profile_id == run.profile_id and resp.profile_version == run.profile_version
    official = s.scalar(select(RequirementPriorityScore).where(RequirementPriorityScore.run_id == m["run"],
                                                               RequirementPriorityScore.requirement_id == m["reqs"]["Touch screen, operation"]))
    assert f"Final score {official.final_score}, rank {official.rank}" in resp.items[0].text


# --- 4-6 agent discussion ------------------------------------------------------------------


def test_agent_discussion_uses_selected_configuration(env):
    s, m, _ = env
    rid = m["reqs"]["Touch screen, operation"]
    resp = agent_chat.discuss_agent(s, m["run"], rid, m["agents"]["demand"])
    text = " ".join(i.text for i in resp.items)
    assert "Log-normalised unique feedback count" in text and "70 distinct G60 comments" in text and "4000" in text
    assert resp.evidence_refs and all(r["table"] == "feedback" for r in resp.evidence_refs)
    custom = agent_chat.discuss_agent(s, m["run"], rid, m["agents"]["feasibility"])
    assessment = [i for i in custom.items if i.kind == "ASSESSMENT"]
    assert assessment and "Synthetic test rationale (feasibility)" in assessment[0].text and "test.engineer" in assessment[0].text
    assert "Technology Feasibility Agent" in custom.title and "Human assessment score" in " ".join(i.text for i in custom.items)


def test_agents_are_discovered_dynamically(env):
    s, m, _ = env
    names = {a["name"] for a in agent_chat.run_agent_options(s, m["run"])}
    assert names == {"Customer Demand Agent", "Evidence Coverage Agent", "Customer Impact Agent", "Technology Feasibility Agent"}


def test_missing_scores_are_explained(env):
    s, m, _ = env
    hud = m["reqs"]["Head-Up Display"]
    resp = da.answer(s, hud, m["run"], intent="missing")
    lims = [i.text for i in resp.items if i.kind == "LIMITATION"]
    assert any("Customer Impact Agent" in l and "No human assessment" in l for l in lims)
    assert any("Technology Feasibility Agent" in l for l in lims)
    chat = agent_chat.discuss_agent(s, m["run"], hud, m["agents"]["impact"])
    assert any("UNSCORED" in i.text for i in chat.items) and any("Missing: a human assessment" in i.text for i in chat.items)
    prio = da.answer(s, hud, m["run"], intent="explain_priority")
    assert "INCOMPLETE" in prio.items[0].text and "no formal rank" in prio.items[0].text


# --- 7-9 debate ------------------------------------------------------------------------------


def test_debate_uses_saved_scores_and_can_abstain(env):
    s, m, _ = env
    a, b, hud = m["reqs"]["Touch screen, operation"], m["reqs"]["Electric Range"], m["reqs"]["Head-Up Display"]
    d = deb.run_debate(s, m["run"], a, b)
    stored_a = {r["agent_id"]: r["score"] for r in dp.agent_results(s, m["run"], a)}
    stored_b = {r["agent_id"]: r["score"] for r in dp.agent_results(s, m["run"], b)}
    for p in d.round1_positions:
        assert p.score_a == stored_a[p.agent_id] and p.score_b == stored_b[p.agent_id]
        expected = ("FAVOURS_A" if p.score_a > p.score_b else "FAVOURS_B" if p.score_b > p.score_a else "NO_DIFFERENCE")
        assert p.position == expected
    assert "A ranks higher" in d.round4_synthesis["ranking_under_profile"]
    feas = next(p for p in d.round1_positions if p.agent_id == m["agents"]["feasibility"])
    assert "Synthetic test rationale (feasibility)" in feas.basis_a
    d2 = deb.run_debate(s, m["run"], a, hud)
    assert {p.position for p in d2.round1_positions if p.agent_id in (m["agents"]["impact"], m["agents"]["feasibility"])} \
        == {"UNABLE_TO_CONCLUDE"}
    assert "No winner is declared" in d2.round4_synthesis["ranking_under_profile"]
    assert d2.round4_synthesis["sensitivity"]["classification"] == "not computable"


def test_debate_sensitivity_is_exact(env):
    s, m, _ = env
    a, b = m["reqs"]["Touch screen, operation"], m["reqs"]["Electric Range"]
    d = deb.run_debate(s, m["run"], a, b)
    for shift in d.round4_synthesis["sensitivity"]["single_agent_shifts_to_tie"]:
        agent_id = next(p.agent_id for p in d.round1_positions if p.agent_name == shift["agent"])
        base = wis.current_weights(s, m["run"])
        target = shift["weight_at_tie"]
        rest = 100 - target
        others = {k: v for k, v in base.items() if k != agent_id}
        weights = {agent_id: target, **{k: rest * v / sum(others.values()) for k, v in others.items()}}
        scores_a = {r["agent_id"]: r["score"] for r in dp.agent_results(s, m["run"], a)}
        scores_b = {r["agent_id"]: r["score"] for r in dp.agent_results(s, m["run"], b)}
        diff = sum(weights[k] * (scores_a[k] - scores_b[k]) for k in weights) / 100
        assert abs(diff) < 0.02


def test_saved_debate_is_linked_to_run(env):
    s, m, _ = env
    d = deb.run_debate(s, m["run"], m["reqs"]["Touch screen, operation"], m["reqs"]["Electric Range"], USER, save=True)
    saved = deb.list_debates(s, run_id=m["run"])
    assert saved[0]["debate_id"] == d.debate_id and saved[0]["run_id"] == m["run"]


# --- 10-15 what-if, scenarios, comparison ------------------------------------------------------------


def test_what_if_requires_100_percent(env):
    s, m, _ = env
    with pytest.raises(wis.SimulationError):
        wis.simulate(s, m["run"], {m["agents"]["demand"]: 60, m["agents"]["coverage"]: 30})
    with pytest.raises(wis.SimulationError):
        wis.simulate(s, m["run"], {m["agents"]["demand"]: 101, m["agents"]["coverage"]: -1})


def test_simulation_recomputes_ranks_without_mutating_official_results(env):
    s, m, _ = env
    before = official_snapshot(s)
    w = {m["agents"]["demand"]: 70, m["agents"]["coverage"]: 30}  # data-only agents: everyone complete
    sim = wis.simulate(s, m["run"], w)
    assert official_snapshot(s) == before
    expected = {}
    for rid in m["reqs"].values():
        res = {r["agent_id"]: r["score"] for r in dp.agent_results(s, m["run"], rid)}
        expected[rid] = round(res[m["agents"]["demand"]] * 0.7 + res[m["agents"]["coverage"]] * 0.3, 2)
    got = {r.requirement_id: r.simulated_final for r in sim.rows}
    assert got == pytest.approx(expected, abs=0.006)
    order = sorted(expected, key=lambda k: -expected[k])
    assert [next(r.simulated_rank for r in sim.rows if r.requirement_id == rid) for rid in order] == [1, 2, 3]
    hud = next(r for r in sim.rows if r.requirement_id == m["reqs"]["Head-Up Display"])
    assert hud.original_rank is None and hud.simulated_rank is not None  # complete under data-only agents
    assert all("SIMULATION" in n or "re-ranked" in n or "Agents" in n for n in sim.notes)


def test_incomplete_stays_unranked_in_simulation(env):
    s, m, _ = env
    sim = wis.simulate(s, m["run"], {m["agents"]["demand"]: 50, m["agents"]["impact"]: 50})
    hud = next(r for r in sim.rows if r.requirement_id == m["reqs"]["Head-Up Display"])
    assert hud.simulated_status == "INCOMPLETE" and hud.simulated_rank is None and hud.simulated_final is None
    assert hud.missing_agents == ["Customer Impact Agent"]


def test_added_agent_without_score_is_incomplete(env):
    s, m, _ = env
    extra, _ = ac.create_agent(s, ac.AgentSpec(name="Strategic Relevance Agent", objective="Strategy",
                                               execution_method="HUMAN_ASSESSMENT", scoring_method="HUMAN_SCORE",
                                               resources=["RES_HUMAN_ASSESSMENTS", "RES_REQUIREMENTS"]), USER)
    sim = wis.simulate(s, m["run"], {m["agents"]["demand"]: 50, extra.agent_id: 50})
    assert all(r.simulated_status == "INCOMPLETE" and r.simulated_rank is None for r in sim.rows)
    assert any("computed now from current data" in c["basis"] for c in sim.rows[0].contributions)


def test_saved_scenario_preserves_configuration_and_leaves_official_data(env):
    s, m, _ = env
    before = official_snapshot(s)
    sim = wis.simulate(s, m["run"], {m["agents"]["demand"]: 70, m["agents"]["coverage"]: 30})
    scn = wis.save_scenario(s, sim, "Demand heavy", USER)
    assert official_snapshot(s) == before
    loaded = wis.load_scenario(s, scn.scenario_id)
    assert loaded.configuration == sim.configuration
    assert [(r.requirement_id, r.simulated_final, r.simulated_rank) for r in loaded.rows] == \
           [(r.requirement_id, r.simulated_final, r.simulated_rank) for r in sim.rows]
    assert s.scalar(select(func.count()).select_from(PrioritizationScenario)) == 1


def test_comparison_discloses_input_differences(env):
    s, m, _ = env
    a = m["reqs"]["Touch screen, operation"]
    weights_only = sc.compare_sides(s, sc.side_from_run(s, m["run"]),
                                    sc.side_from_simulation(s, wis.simulate(s, m["run"], {m["agents"]["demand"]: 40,
                                                                                         m["agents"]["coverage"]: 10,
                                                                                         m["agents"]["impact"]: 25,
                                                                                         m["agents"]["feasibility"]: 25})))
    assert "attributable to weight changes" in weights_only["attribution"] and not weights_only["input_differences"]
    dp.record_human_assessment(s, a, m["agents"]["impact"], 30, "Synthetic re-assessment", "test.assessor")
    rerun = dp.run_prioritization(s, m["profile"], USER)
    changed = sc.compare_sides(s, sc.side_from_run(s, m["run"]), sc.side_from_run(s, rerun.run_id))
    assert "cannot be attributed to weights alone" in changed["attribution"]
    assert any("Customer Impact Agent" in d and "80.0 vs 30.0" in d for d in changed["input_differences"])
    fewer = sc.compare_sides(s, sc.side_from_run(s, m["run"]),
                             sc.side_from_simulation(s, wis.simulate(s, m["run"], {m["agents"]["demand"]: 100})))
    assert "different agent set" in fewer["attribution"]


# --- 16-17 challenges, history ----------------------------------------------------------------


def test_challenge_workflow_preserves_history_and_scores(env):
    s, m, _ = env
    rid = m["reqs"]["Electric Range"]
    before = official_snapshot(s)
    c = pc.create_challenge(s, rid, m["run"], "Range anxiety is more serious than this impact score suggests.", USER,
                            "SCORE", m["agents"]["impact"])
    view = pc.challenge_view(s, c.challenge_id)
    assert view["status"] == "OPEN" and view["profile_version"] == 1
    assert any("Human assessment" in n for n in view["needs_review"])
    assert view["evidence_snapshot"][0]["score"] == 70.0
    with pytest.raises(pc.ChallengeError):
        pc.respond(s, c.challenge_id, "resolve", "test.lead", "")
    pc.respond(s, c.challenge_id, "start_review", "test.lead", "Reviewing")
    with pytest.raises(pc.ChallengeError):
        pc.respond(s, c.challenge_id, "resolve", "test.lead", "ok", linked_assessment_id="ASM-NOPE")
    new = dp.record_human_assessment(s, rid, m["agents"]["impact"], 85, "Synthetic reassessment after review", "test.assessor")
    pc.respond(s, c.challenge_id, "resolve", "test.lead", "Reassessed via Part 5", linked_assessment_id=new.assessment_id)
    view = pc.challenge_view(s, c.challenge_id)
    assert [e["action"] for e in view["events"]] == ["create", "start_review", "resolve"]
    assert view["linked_assessment_id"] == new.assessment_id
    after = official_snapshot(s)
    assert after[0] == before[0] and after[1] == before[1]  # official run results untouched by the challenge
    assert rr.get_requirement(s, rid).review_status == "DRAFT"


def test_chat_history_stays_linked_to_original_run(env):
    s, m, _ = env
    rid = m["reqs"]["Touch screen, operation"]
    conv, _ = da.ask(s, rid, m["run"], USER, intent="explain_priority")
    da.ask(s, rid, m["run"], USER, intent="evidence", conversation_id=conv)
    msgs = s.scalars(select(DecisionMessage).where(DecisionMessage.conversation_id == conv)).all()
    assert len(msgs) == 2 and {x.run_id for x in msgs} == {m["run"]}
    assert all("customer_feedback" not in x.response_json for x in msgs)
    dp.run_prioritization(s, m["profile"], USER)  # a newer run appears
    hist = da.conversation_history(s, conv)
    assert hist["run_id"] == m["run"] and any("Historical context" in n for n in hist["notices"])
    assert hist["messages"][0]["response"].run_id == m["run"]


# --- 18-23 reports ----------------------------------------------------------------------------------


def test_excel_report(env):
    s, m, _ = env
    content, name = mr.export_report(s, m["run"], "xlsx", USER)
    wb = load_workbook(io.BytesIO(content))
    assert {"Overview", "Requirements", "Agent scores", "Weights", "Evidence", "Critical flags"} <= set(wb.sheetnames)
    ws = wb["Requirements"]
    assert ws.auto_filter.ref and ws.freeze_panes == "A2"
    rows = list(ws.iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    hud = next(r for r in body if r[0] == m["reqs"]["Head-Up Display"])
    assert hud[header.index("priority_status")] == "INCOMPLETE"
    assert hud[header.index("rank")] in (None, "") and hud[header.index("final_score")] in (None, "")
    overview = dict(list(wb["Overview"].iter_rows(values_only=True))[1:])
    assert "NOT an official BMW" in overview["Methodology"] and m["run"] in overview["Prioritization run"]
    assert name.endswith(".xlsx") and s.scalar(select(func.count()).select_from(ManagementReportExport)) == 1


def test_csv_report_preserves_references(env):
    s, m, _ = env
    content, _ = mr.export_report(s, m["run"], "csv", USER)
    rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    assert len(rows) == 3 * 4  # requirements x agents
    demand = [r for r in rows if r["agent"] == "Customer Demand Agent"]
    assert all(r["evidence_references"].startswith("G60-") for r in demand)
    for r in rows:
        if r["priority_status"] == "INCOMPLETE":
            assert r["rank"] == "" and r["final_score"] == ""
        if r["status"] == "UNSCORED":
            assert r["score"] == "" and r["weighted_contribution"] == ""
    assert all(r["run_id"] == m["run"] for r in rows)


def test_pdf_report(env):
    s, m, _ = env
    rid = m["reqs"]["Touch screen, operation"]
    ci_flag = __import__("services.critical_issues", fromlist=["x"])
    key = s.scalar(select(__import__("core.requirement_models", fromlist=["x"]).RequirementEvidenceLink.source_key)
                   .where(__import__("core.requirement_models", fromlist=["x"]).RequirementEvidenceLink.requirement_id == rid))
    f = ci_flag.raise_flag(s, m["reqs"]["Head-Up Display"], "SAFETY", "Synthetic test flag", USER)
    ci_flag.decide_flag(s, f.flag_id, "CONFIRMED", "test.lead", "Confirmed in synthetic test")
    content, name = mr.export_report(s, m["run"], "pdf", USER)
    doc = pymupdf.open("pdf", content)
    text = " ".join(p.get_text() for p in doc)
    assert doc.page_count >= 2 and text.count("prototype methodology") >= doc.page_count
    assert "Incomplete requirements - not ranked" in text and "Confirmed critical issues" in text
    assert m["run"] in text and key


def test_reports_exclude_verbatim_and_secrets_by_default(env, monkeypatch):
    s, m, _ = env
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-SECRET-TEST-VALUE")
    rid = m["reqs"]["Touch screen, operation"]
    one_comment = s.scalar(select(__import__("core.models", fromlist=["x"]).Feedback.customer_feedback)
                           .where(__import__("core.models", fromlist=["x"]).Feedback.original_feedback_id ==
                                  dp.agent_results(s, m["run"], rid)[0]["source_refs"][0]["key"]))
    for fmt in ("xlsx", "csv", "pdf"):
        content, _ = mr.export_report(s, m["run"], fmt, USER)
        blob = content if fmt != "xlsx" else b"".join(
            str(c.value).encode() for ws in load_workbook(io.BytesIO(content)).worksheets for row in ws.iter_rows() for c in row)
        assert b"sk-ant-SECRET" not in blob
        if fmt != "pdf":
            assert one_comment[:40].encode() not in blob
    content, _ = mr.export_report(s, m["run"], "xlsx", USER, mr.ReportOptions(include_verbatim_feedback=True))
    ev = list(load_workbook(io.BytesIO(content))["Evidence"].iter_rows(values_only=True))
    assert any(r[-1] for r in ev[1:])  # verbatim only when explicitly selected


def test_report_includes_selected_scenarios_and_challenges(env):
    s, m, _ = env
    scn = wis.save_scenario(s, wis.simulate(s, m["run"], {m["agents"]["demand"]: 70, m["agents"]["coverage"]: 30}), "S", USER)
    pc.create_challenge(s, m["reqs"]["Electric Range"], m["run"], "Synthetic test challenge", USER)
    data = mr.build_report_data(s, m["run"], mr.ReportOptions(scenario_ids=[scn.scenario_id], include_challenges=True))
    assert len(data["tables"]["Scenarios"]) == 3 and all(r["label"].startswith("SIMULATION") for r in data["tables"]["Scenarios"])
    assert data["tables"]["Challenges"][0]["argument"] == "Synthetic test challenge"


# --- 23 critical flags, LLM gate -----------------------------------------------------------------------


def test_critical_flags_visible_in_assistant(env):
    s, m, _ = env
    from services import critical_issues as ci

    hud = m["reqs"]["Head-Up Display"]
    f = ci.raise_flag(s, hud, "SEVERE_RELIABILITY", "Synthetic test flag", USER)
    ci.decide_flag(s, f.flag_id, "CONFIRMED", "test.lead", "Confirmed in synthetic test")
    resp = da.answer(s, hud, m["run"], intent="challenge")
    assert any("SEVERE_RELIABILITY (CONFIRMED)" in i.text for i in resp.items)
    sim = wis.simulate(s, m["run"], {m["agents"]["demand"]: 100})
    assert next(r for r in sim.rows if r.requirement_id == hud).critical_flags[0]["status"] == "CONFIRMED"


def test_llm_mode_disabled_and_validated(env):
    s, m, _ = env
    status = dl.llm_mode_status()
    assert not status.enabled and any("not authorised" in r for r in status.reasons)
    resp = da.answer(s, m["reqs"]["Touch screen, operation"], m["run"], intent="explain_priority")
    with pytest.raises(PermissionError):
        dl.rephrase(resp, confirmed=True, cost_acknowledged=dl.estimate_cost(resp))

    class FakeLocal:
        name, is_local = "fake-local", True

        def __init__(self, text):
            self.text = text

        def rephrase(self, system, payload, max_tokens):
            return self.text

    with pytest.raises(PermissionError):
        dl.rephrase(resp, confirmed=False, cost_acknowledged=None, model=FakeLocal("x"))
    bad = dl.rephrase(resp, confirmed=True, cost_acknowledged={}, model=FakeLocal("It scored 99.9 because of G60-9999."))
    assert bad["text"] is None and bad["errors"]
    first = resp.items[0].text
    good = dl.rephrase(resp, confirmed=True, cost_acknowledged={}, model=FakeLocal(first))
    assert good["text"] == first and not good["errors"]


# --- 24-25 UI -----------------------------------------------------------------------------------------


def _app(db, monkeypatch):
    from streamlit.testing.v1 import AppTest
    import streamlit as st

    monkeypatch.setenv("PRODUCTMIND_DB_PATH", str(db))
    st.cache_data.clear()
    st.cache_resource.clear()
    return AppTest.from_file(str(ROOT / "app.py"), default_timeout=180)


def test_empty_state(base_db, tmp_path, monkeypatch):
    path = tmp_path / "empty.db"
    shutil.copy2(base_db, path)
    app = _app(path, monkeypatch)
    app.switch_page("ui/pages/decision_intelligence.py").run()
    assert not app.exception and any("No saved requirements yet" in i.value for i in app.info)


def test_decision_page_flows(env, monkeypatch):
    s, m, path = env
    s.commit()
    app = _app(path, monkeypatch)
    app.switch_page("ui/pages/decision_intelligence.py").run()
    assert not app.exception
    for label in ("Explain Ranking", "Show Evidence", "Challenge Priority", "Start Debate", "Simulate", "Export PDF",
                  "Export Excel", "Export CSV"):
        next(b for b in app.button if b.label == label).click().run()
        assert not app.exception, (label, [e.value for e in app.exception])
    assert {b.label for b in app.get("download_button")} >= {"Download PDF", "Download XLSX", "Download CSV"}
    html = " ".join(x.value for x in app.markdown)
    assert "Evidence-based explanation — not AI-generated" in html


def test_all_pages_still_load(env, monkeypatch):
    s, m, path = env
    s.commit()
    app = _app(path, monkeypatch)
    app.run()
    for page in ("ui/pages/evidence_explorer.py", "ui/pages/customer_insights.py", "ui/pages/vehicle_requirements.py",
                 "ui/pages/agent_prioritization.py", "ui/pages/decision_intelligence.py", "ui/pages/data_quality.py"):
        app.switch_page(page).run()
        assert not app.exception, (page, [e.value for e in app.exception])


def test_tied_top_contributors_are_reported_as_ties(env):
    s, m, _ = env
    # Touch screen: coverage 100 x 20% = 20.0 and impact 80 x 25% = 20.0 (synthetic) -> a tie
    resp = da.answer(s, m["reqs"]["Touch screen, operation"], m["run"], intent="explain_priority")
    tie = [i.text for i in resp.items if "Largest weighted contribution" in i.text]
    assert tie and "shared" in tie[0] and "Evidence Coverage Agent" in tie[0] and "Customer Impact Agent" in tie[0]
