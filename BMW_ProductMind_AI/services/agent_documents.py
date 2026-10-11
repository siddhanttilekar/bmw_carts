"""Agent resource uploads (ProductMind AI 2.0).

Uploaded files are UNTRUSTED content. They are parsed to plain text only:

* PDF  - PyMuPDF text extraction (no JavaScript, no embedded files)
* XLSX - openpyxl ``read_only`` + ``data_only`` (cached values; formulas and
         macros are never evaluated)
* CSV / TXT / MD - decoded text

Text is stored as chunks in ``agent_document_chunks`` (no file is written to
disk) and is only ever searched with keyword matching. Instructions inside a
document are never executed or followed. Retrieval is scoped to the owning
agent - one agent can never read another agent's uploads.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import uuid
from datetime import datetime, timezone

from sqlalchemy import select
from sqlalchemy.orm import Session

from core.evidence_models import AgentDocument, AgentDocumentChunk

ALLOWED_TYPES = ("pdf", "xlsx", "csv", "txt", "md")
MAX_BYTES = 10 * 1024 * 1024
MAX_CHUNKS = 2000
CHUNK_CHARS = 1500
STOPWORDS = frozenset("""the and for with that this from shall should must will have has into over under their there
about than when which while where what your more less next generation bmw series vehicle customer customers need needs
requirement""".split())


class DocumentError(ValueError):
    pass


def _now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _text_chunks(text: str, prefix: str) -> list[tuple[str, str]]:
    out, buf, start, line_no = [], [], 1, 0
    for line_no, line in enumerate(text.splitlines(), start=1):
        buf.append(line)
        if sum(len(b) for b in buf) >= CHUNK_CHARS:
            out.append((f"{prefix}lines {start}-{line_no}", "\n".join(buf)))
            buf, start = [], line_no + 1
    if any(b.strip() for b in buf):
        out.append((f"{prefix}lines {start}-{line_no}", "\n".join(buf)))
    return out


def parse_document(file_name: str, content: bytes) -> list[tuple[str, str]]:
    """Returns [(location, text)]. Raises DocumentError for unsupported or
    unreadable files."""
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext not in ALLOWED_TYPES:
        raise DocumentError(f"Unsupported file type '.{ext}'. Allowed: {', '.join(ALLOWED_TYPES).upper()}.")
    if not content:
        raise DocumentError("The file is empty.")
    if len(content) > MAX_BYTES:
        raise DocumentError(f"The file exceeds {MAX_BYTES // (1024 * 1024)} MB.")
    try:
        if ext == "pdf":
            import fitz

            chunks = []
            with fitz.open(stream=content, filetype="pdf") as doc:
                for i, page in enumerate(doc, start=1):
                    text = page.get_text("text")
                    chunks += _text_chunks(text, f"page {i}, ")
            return chunks
        if ext == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            chunks = []
            for ws in wb.worksheets:
                rows, first = [], 1
                for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    cells = [str(v) for v in row if v is not None and str(v).strip()]
                    if cells:
                        rows.append(" | ".join(cells))
                    if len(rows) >= 25:
                        chunks.append((f"sheet {ws.title}, rows {first}-{r}", "\n".join(rows)))
                        rows, first = [], r + 1
                if rows:
                    chunks.append((f"sheet {ws.title}, rows {first}-{ws.max_row or first}", "\n".join(rows)))
            wb.close()
            return chunks
        text = content.decode("utf-8-sig", errors="replace")
        if ext == "csv":
            lines = [" | ".join(c for c in row if c.strip()) for row in csv.reader(io.StringIO(text))]
            text = "\n".join(line for line in lines if line)
        return _text_chunks(text, "")
    except DocumentError:
        raise
    except Exception as exc:  # noqa: BLE001 - any parser failure is a user-facing rejection
        raise DocumentError(f"The file could not be read as {ext.upper()}: {type(exc).__name__}.") from exc


def upload_document(session: Session, agent_id: str, file_name: str, content: bytes, uploaded_by: str) -> AgentDocument:
    from services.agent_configuration import get_agent

    if not (uploaded_by or "").strip():
        raise DocumentError("A user identifier is required.")
    agent = get_agent(session, agent_id)
    if agent.deleted:
        raise DocumentError("Documents cannot be added to a removed agent.")
    safe_name = re.sub(r"[^\w.\- ]", "_", file_name.split("/")[-1].split("\\")[-1])[:200] or "upload"
    chunks = parse_document(safe_name, content)
    if not chunks:
        raise DocumentError("No readable text was found in the file.")
    if len(chunks) > MAX_CHUNKS:
        raise DocumentError(f"The document is too large ({len(chunks)} passages; limit {MAX_CHUNKS}).")
    digest = hashlib.sha256(content).hexdigest()
    existing = session.scalar(select(AgentDocument).where(AgentDocument.agent_id == agent_id, AgentDocument.sha256 == digest))
    if existing is not None:
        if not existing.active:
            existing.active = True
        return existing
    doc = AgentDocument(document_id=f"DOC-{uuid.uuid4().hex[:10].upper()}", agent_id=agent_id, file_name=safe_name,
                        file_type=safe_name.rsplit(".", 1)[-1].lower(), sha256=digest, size_bytes=len(content),
                        chunk_count=len(chunks), active=True, uploaded_by=uploaded_by.strip(), created_at=_now())
    session.add(doc)
    session.flush()
    session.add_all(AgentDocumentChunk(chunk_id=f"{doc.document_id}-{i:04d}", document_id=doc.document_id, chunk_index=i,
                                       location=loc[:128], text=text) for i, (loc, text) in enumerate(chunks))
    session.flush()
    return doc


def set_document_active(session: Session, agent_id: str, document_id: str, active: bool) -> AgentDocument:
    doc = session.get(AgentDocument, document_id)
    if doc is None or doc.agent_id != agent_id:
        raise DocumentError("Unknown document for this agent.")
    doc.active = active
    return doc


def agent_documents(session: Session, agent_id: str, active_only: bool = True) -> list[dict]:
    stmt = select(AgentDocument).where(AgentDocument.agent_id == agent_id)
    if active_only:
        stmt = stmt.where(AgentDocument.active.is_(True))
    return [{"document_id": d.document_id, "file_name": d.file_name, "file_type": d.file_type, "sha256": d.sha256,
             "size_bytes": d.size_bytes, "passages": d.chunk_count, "active": d.active, "uploaded_by": d.uploaded_by,
             "created_at": d.created_at} for d in session.scalars(stmt.order_by(AgentDocument.created_at))]


def keywords(*texts: str) -> list[str]:
    words = re.findall(r"[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ\-]{3,}", " ".join(t or "" for t in texts).lower())
    return sorted({w for w in words if w not in STOPWORDS})


def search_agent_documents(session: Session, agent_id: str, terms: list[str], min_hits: int = 2) -> tuple[int, list[dict]]:
    """Keyword retrieval restricted to ``agent_id``'s active documents.
    Returns (total passages searched, matching passages)."""
    chunks = list(session.scalars(
        select(AgentDocumentChunk).join(AgentDocument, AgentDocument.document_id == AgentDocumentChunk.document_id)
        .where(AgentDocument.agent_id == agent_id, AgentDocument.active.is_(True))
        .order_by(AgentDocumentChunk.chunk_id)))
    needed = min(min_hits, len(terms)) if terms else 1
    hits = []
    for c in chunks:
        low = c.text.lower()
        found = [t for t in terms if re.search(rf"\b{re.escape(t)}\b", low)]
        if terms and len(found) >= needed:
            hits.append({"chunk_id": c.chunk_id, "document_id": c.document_id, "location": c.location,
                         "matched_terms": found, "excerpt": c.text[:300]})
    return len(chunks), hits
