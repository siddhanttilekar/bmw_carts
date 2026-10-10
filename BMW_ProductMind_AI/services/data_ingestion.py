"""Data ingestion: supplied Excel / PDF / PPTX files -> SQLite.

Principles (see README):
* Original files are the single source of truth. The consolidated workbook is
  compared against them and recorded in ``dataset_overlaps``; its copies are
  not imported a second time.
* Values are stored verbatim. Missing values stay NULL. Nothing is invented.
* Every record carries provenance (file, sheet/page, row, record_id, timestamp).
* Internal BMW feedback and public evidence live in separate tables.
"""
from __future__ import annotations

import hashlib
import json
import logging
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
import pymupdf
from openpyxl.utils import get_column_letter
from pptx import Presentation
from sqlalchemy import insert
from sqlalchemy.orm import Session

from core import config
from core.models import (
    CompetitorEvidence,
    CustomerStudy,
    DatasetNote,
    DatasetOverlap,
    Equipment,
    EquipmentAvailability,
    Feedback,
    FeedbackClassification,
    PossibleDuplicate,
    PublicFeedback,
    SalesVolume,
    SourceDocument,
    SourceDocumentPage,
)

log = logging.getLogger(__name__)

RECORD_NAMESPACE = uuid.UUID("6f1d3c9e-2b7a-4e1f-9a51-5d0c8b2e7a10")


class IngestionError(Exception):
    """Raised when a file does not match the structure the parser expects."""


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------


def utc_now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def make_record_id(*parts: Any) -> str:
    """Deterministic ID from provenance, so re-imports produce identical IDs."""
    return str(uuid.uuid5(RECORD_NAMESPACE, "|".join(str(p) for p in parts)))


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean(value: Any) -> Any:
    """Convert pandas missing markers to None; otherwise return value unchanged."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def clean_str(value: Any) -> str | None:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value)


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def bulk_insert(session: Session, model: type, rows: list[dict]) -> int:
    if rows:
        session.execute(insert(model), rows)
    return len(rows)


@dataclass
class IngestionContext:
    session: Session
    imported_at: datetime = field(default_factory=utc_now)
    data_dir: Path | None = None
    # canonical file name -> manifest entry (see services/source_discovery.py)
    manifest: dict = field(default_factory=dict)


def register_document(
    ctx: IngestionContext,
    path: Path,
    data_origin: str,
    is_source_of_truth: bool,
    sheet_or_page_count: int | None,
    status: str = "imported",
    notes: str | None = None,
) -> SourceDocument:
    from services.source_discovery import identify_file  # local import: avoids a cycle

    rel = path.relative_to(ctx.data_dir) if ctx.data_dir else path.name
    origin_info = ctx.manifest.get(path.name, {})
    doc = SourceDocument(
        document_id=make_record_id("document", path.name),
        file_name=path.name,
        relative_path=str(rel),
        original_file_names=json.dumps(origin_info.get("original_file_names", [path.name])),
        original_source_paths=json.dumps(origin_info["original_source_paths"]) if origin_info else None,
        dataset_type=identify_file(path).dataset_type,
        file_type=path.suffix.lower().lstrip("."),
        data_origin=data_origin,
        sha256=file_sha256(path),
        size_bytes=path.stat().st_size,
        sheet_or_page_count=sheet_or_page_count,
        is_source_of_truth=is_source_of_truth,
        ingestion_status=status,
        ingestion_notes=notes,
        imported_at=ctx.imported_at,
    )
    ctx.session.add(doc)
    ctx.session.flush()
    return doc


def provenance(
    ctx: IngestionContext,
    doc: SourceDocument,
    *,
    sheet: str | None = None,
    page: int | None = None,
    row: int | None = None,
    key: Iterable[Any] = (),
) -> dict:
    return {
        "record_id": make_record_id(doc.file_name, sheet, page, row, *key),
        "source_document_id": doc.document_id,
        "source_file": doc.file_name,
        "source_sheet": sheet,
        "source_page": page,
        "source_row": row,
        "imported_at": ctx.imported_at,
    }


def read_sheet(path: Path, sheet: str, header: int | None = 0) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, header=header, dtype=object)


def excel_sheet_names(path: Path) -> list[str]:
    with pd.ExcelFile(path) as workbook:
        return list(workbook.sheet_names)


def require_columns(df: pd.DataFrame, expected: list[str], where: str) -> None:
    actual = [str(c) for c in df.columns]
    missing = [c for c in expected if c not in actual]
    if missing:
        raise IngestionError(f"{where}: missing expected columns {missing}; found {actual}")


# ---------------------------------------------------------------------------
# Key/value note sheets (README, Method and caveats, START HERE)
# ---------------------------------------------------------------------------


def ingest_note_sheet(ctx: IngestionContext, doc: SourceDocument, path: Path, sheet: str) -> int:
    df = read_sheet(path, sheet, header=None)
    rows = []
    for idx, rec in df.iterrows():
        values = [clean_str(v) for v in rec.tolist()]
        if not any(values):
            continue
        rows.append(
            {
                **provenance(ctx, doc, sheet=sheet, row=idx + 1),
                "note_key": values[0],
                "note_value": " | ".join(v for v in values[1:] if v) or None,
            }
        )
    return bulk_insert(ctx.session, DatasetNote, rows)


# ---------------------------------------------------------------------------
# Internal BMW feedback
# ---------------------------------------------------------------------------


def classification_status(vfc_level2: str | None) -> str:
    if vfc_level2 is None:
        return "missing"
    if vfc_level2 == "no_class_found":
        return "no_class_found"
    return "classified"


def ingest_feedback_file(ctx: IngestionContext, doc: SourceDocument, path: Path, vehicle_code: str) -> dict:
    sheet = config.FEEDBACK_SHEET
    df = read_sheet(path, sheet)
    require_columns(df, config.FEEDBACK_COLUMNS, f"{path.name}/{sheet}")

    derivates = set(df["Derivate (e-code)"].dropna().astype(str))
    if derivates != {vehicle_code}:
        raise IngestionError(f"{path.name}: expected e-code {vehicle_code}, found {sorted(derivates)}")

    comment_fields = ["Source", "Brand", "Sales Description", "Country", "Engine Type", "Customer Feedback"]
    inconsistent = [
        c for c in comment_fields if (df.groupby("ID")[c].nunique(dropna=False) > 1).any()
    ]

    feedback_rows: dict[str, dict] = {}
    class_rows: list[dict] = []
    for idx, rec in df.iterrows():
        excel_row = idx + 2  # header occupies row 1
        fid = clean_str(rec["ID"])
        if fid is None:
            raise IngestionError(f"{path.name} row {excel_row}: missing feedback ID")
        if fid not in feedback_rows:
            feedback_rows[fid] = {
                **provenance(ctx, doc, sheet=sheet, row=excel_row, key=("feedback", fid)),
                "data_origin": "internal_bmw",
                "original_feedback_id": fid,
                "vehicle_code": vehicle_code,
                "feedback_source": clean_str(rec["Source"]),
                "brand": clean_str(rec["Brand"]),
                "sales_description": clean_str(rec["Sales Description"]),
                "country": clean_str(rec["Country"]),
                "engine_type": clean_str(rec["Engine Type"]),
                "customer_feedback": clean_str(rec["Customer Feedback"]) or "",
                "classification_row_count": 0,
            }
        feedback_rows[fid]["classification_row_count"] += 1
        level2 = clean_str(rec["Vfc level2 Name"])
        class_rows.append(
            {
                **provenance(ctx, doc, sheet=sheet, row=excel_row, key=("classification",)),
                "feedback_record_id": feedback_rows[fid]["record_id"],
                "original_feedback_id": fid,
                "vehicle_code": vehicle_code,
                "feedback_type": clean_str(rec["Feedback Type"]),
                "vfc_level2_3_name": clean_str(rec["Vfc level2 3 Name"]),
                "vfc_level2_name": level2,
                "vfc_level3_name": clean_str(rec["Vfc level3 Name"]),
                "derived_classification_status": classification_status(level2),
            }
        )

    n_feedback = bulk_insert(ctx.session, Feedback, list(feedback_rows.values()))
    n_class = bulk_insert(ctx.session, FeedbackClassification, class_rows)
    flag_feedback_text_duplicates(ctx, list(feedback_rows.values()))
    return {
        "source_rows": len(df),
        "feedback": n_feedback,
        "classifications": n_class,
        "inconsistent_fields_within_id": inconsistent,
    }


def flag_feedback_text_duplicates(ctx: IngestionContext, feedback_rows: list[dict]) -> None:
    """Flag different feedback IDs with identical normalised text. Not merged."""
    groups: dict[str, list[dict]] = {}
    for row in feedback_rows:
        groups.setdefault(normalize_text(row["customer_feedback"]), []).append(row)
    dupes = [
        {
            "table_name": "feedback",
            "rule": "identical_normalized_text_different_ids",
            "group_key": text[:500],
            "record_ids": ",".join(r["record_id"] for r in members),
            "original_ids": ",".join(r["original_feedback_id"] for r in members),
            "record_count": len(members),
            "review_status": "unreviewed",
        }
        for text, members in groups.items()
        if len(members) > 1
    ]
    bulk_insert(ctx.session, PossibleDuplicate, dupes)


# ---------------------------------------------------------------------------
# Public evidence (kept separate from BMW feedback)
# ---------------------------------------------------------------------------

PUBLIC_FEEDBACK_COLUMNS = [
    "ID", "Model", "Model year", "Review date", "Source", "Source type", "Market", "Sentiment",
    "Theme", "Paraphrased owner feedback", "Product implication", "Source URL", "Evidence note",
]
RELATION_RE = re.compile(r"Same (review|owner) as (P\d+)", re.IGNORECASE)


def ingest_public_feedback(ctx: IngestionContext, doc: SourceDocument, path: Path) -> dict:
    sheet = "Public G60 Feedback"
    df = read_sheet(path, sheet)
    require_columns(df, PUBLIC_FEEDBACK_COLUMNS, f"{path.name}/{sheet}")
    rows = []
    for idx, rec in df.iterrows():
        note = clean_str(rec["Evidence note"])
        match = RELATION_RE.search(note or "")
        rows.append(
            {
                **provenance(ctx, doc, sheet=sheet, row=idx + 2),
                "data_origin": "public_external",
                "original_public_id": clean_str(rec["ID"]),
                "model": clean_str(rec["Model"]),
                "model_year": clean_str(rec["Model year"]),
                "review_date": clean_str(rec["Review date"]),
                "source_name": clean_str(rec["Source"]),
                "source_type": clean_str(rec["Source type"]),
                "market": clean_str(rec["Market"]),
                "sentiment": clean_str(rec["Sentiment"]),
                "theme": clean_str(rec["Theme"]),
                "paraphrased_feedback": clean_str(rec["Paraphrased owner feedback"]),
                "product_implication": clean_str(rec["Product implication"]),
                "source_url": clean_str(rec["Source URL"]),
                "evidence_note": note,
                "derived_related_public_id": match.group(2) if match else None,
                "derived_relation_type": f"same_{match.group(1).lower()}" if match else None,
            }
        )
    n = bulk_insert(ctx.session, PublicFeedback, rows)
    n_notes = ingest_note_sheet(ctx, doc, path, "README")
    return {"source_rows": len(df), "public_feedback": n, "notes": n_notes}


COMPETITOR_COLUMNS = [
    "Brand", "Model", "Source type", "Date or period", "Market", "Theme", "Sentiment",
    "Extracted insight (paraphrased)", "Evidence limitation", "Source URL",
]


def ingest_competitor_evidence(ctx: IngestionContext, doc: SourceDocument, path: Path) -> dict:
    sheet = "Public Evidence"
    df = read_sheet(path, sheet)
    require_columns(df, COMPETITOR_COLUMNS, f"{path.name}/{sheet}")
    rows = []
    for idx, rec in df.iterrows():
        brand = clean_str(rec["Brand"])
        rows.append(
            {
                **provenance(ctx, doc, sheet=sheet, row=idx + 2),
                "data_origin": "public_external",
                "brand": brand,
                "model": clean_str(rec["Model"]),
                "source_type": clean_str(rec["Source type"]),
                "date_or_period": clean_str(rec["Date or period"]),
                "market": clean_str(rec["Market"]),
                "theme": clean_str(rec["Theme"]),
                "sentiment": clean_str(rec["Sentiment"]),
                "extracted_insight": clean_str(rec["Extracted insight (paraphrased)"]),
                "evidence_limitation": clean_str(rec["Evidence limitation"]),
                "source_url": clean_str(rec["Source URL"]),
                "derived_is_bmw": (brand or "").strip().upper() == "BMW",
            }
        )
    n = bulk_insert(ctx.session, CompetitorEvidence, rows)
    n_notes = ingest_note_sheet(ctx, doc, path, "Method and caveats")
    return {"source_rows": len(df), "competitor_evidence": n, "notes": n_notes}


# ---------------------------------------------------------------------------
# Customer studies (survey cross-tabs)
# ---------------------------------------------------------------------------

VEHICLE_CODE_RE = re.compile(r"\b([A-Z]\d{2})\b")


def vehicle_code_from_label(label: str) -> str | None:
    match = VEHICLE_CODE_RE.search(label)
    return match.group(1) if match else None


def _is_number(value: Any) -> bool:
    return isinstance(clean(value), (int, float)) and not isinstance(value, bool)


def parse_cn_eu_study(df: pd.DataFrame) -> list[dict]:
    """Layout: row 2 = market (merged across columns), row 3 = vehicle labels and
    the first attribute name in column A; then attribute rows followed by 'Mean'."""
    markets: list[str | None] = []
    current = None
    for col in range(1, df.shape[1]):
        current = clean_str(df.iat[1, col]) or current
        markets.append(current)
    vehicles = {col: clean_str(df.iat[2, col]) for col in range(1, df.shape[1])}
    if not all(vehicles.values()) or not all(markets):
        raise IngestionError("CN_EU_2025: could not read market/vehicle header rows")

    cells, attribute = [], clean_str(df.iat[2, 0])
    for r in range(3, df.shape[0]):
        label = clean_str(df.iat[r, 0])
        values = [df.iat[r, c] for c in range(1, df.shape[1])]
        if label is None and all(clean(v) is None for v in values):
            continue
        if all(clean(v) is None for v in values):
            attribute = label
            continue
        for c in range(1, df.shape[1]):
            if not _is_number(df.iat[r, c]):
                raise IngestionError(f"CN_EU_2025 row {r + 1}: non-numeric value {df.iat[r, c]!r}")
            cells.append(
                {
                    "excel_row": r + 1,
                    "column": get_column_letter(c + 1),
                    "survey_market": markets[c - 1],
                    "vehicle_label": vehicles[c],
                    "attribute": attribute,
                    "measure": label,
                    "measure_base": None,
                    "value_type": "mean_score" if label == "Mean" else "other",
                    "value": float(df.iat[r, c]),
                }
            )
    return cells


def parse_us_study(df: pd.DataFrame) -> list[dict]:
    """Layout: attribute header row (col B empty) followed by 'Sample total' and
    seven answer rows, each '% of column, sample base'. Row 1 carries vehicle labels."""
    vehicles = {c: clean_str(df.iat[0, c]) for c in range(2, df.shape[1])}
    if not all(vehicles.values()):
        raise IngestionError("US_2025: could not read vehicle header row")
    cells, attribute = [], None
    for r in range(df.shape[0]):
        label, base = clean_str(df.iat[r, 0]), clean_str(df.iat[r, 1])
        if label is None and base is None:
            continue  # blank spacer row inside a block
        if base is None:
            attribute = label
            continue
        if attribute is None:
            raise IngestionError(f"US_2025 row {r + 1}: answer row before any attribute")
        for c in range(2, df.shape[1]):
            if not _is_number(df.iat[r, c]):
                raise IngestionError(f"US_2025 row {r + 1}: non-numeric value {df.iat[r, c]!r}")
            cells.append(
                {
                    "excel_row": r + 1,
                    "column": get_column_letter(c + 1),
                    "survey_market": "US",
                    "vehicle_label": vehicles[c],
                    "attribute": attribute,
                    "measure": label,
                    "measure_base": base,
                    "value_type": "column_proportion",
                    "value": float(df.iat[r, c]),
                }
            )
    return cells


def ingest_customer_studies(ctx: IngestionContext, doc: SourceDocument, path: Path) -> dict:
    parsers = {"CN_EU_2025": parse_cn_eu_study, "US_2025": parse_us_study}
    result = {}
    for sheet, parser in parsers.items():
        cells = parser(read_sheet(path, sheet, header=None))
        rows = [
            {
                **provenance(ctx, doc, sheet=sheet, row=cell["excel_row"], key=(cell["column"],)),
                "study_name": sheet,
                "survey_market": cell["survey_market"],
                "vehicle_label": cell["vehicle_label"],
                "vehicle_code": vehicle_code_from_label(cell["vehicle_label"]),
                "attribute": cell["attribute"],
                "measure": cell["measure"],
                "measure_base": cell["measure_base"],
                "value_type": cell["value_type"],
                "value": cell["value"],
                "source_column": cell["column"],
            }
            for cell in cells
        ]
        result[sheet] = bulk_insert(ctx.session, CustomerStudy, rows)
    return result


# ---------------------------------------------------------------------------
# Sales volumes
# ---------------------------------------------------------------------------

VOLUME_COL_RE = re.compile(r"^volume_(\d{4})$")


def ingest_sales(ctx: IngestionContext, doc: SourceDocument, path: Path) -> dict:
    result = {}
    for sheet in excel_sheet_names(path):
        if sheet not in config.SALES_SHEET_VEHICLES:
            raise IngestionError(f"{path.name}: unexpected sales sheet {sheet!r}")
        df = read_sheet(path, sheet)
        require_columns(df, ["market", "market_code"], f"{path.name}/{sheet}")
        year_cols = {c: int(m.group(1)) for c in df.columns if (m := VOLUME_COL_RE.match(str(c)))}
        if not year_cols:
            raise IngestionError(f"{path.name}/{sheet}: no volume_YYYY columns")
        vehicles = config.SALES_SHEET_VEHICLES[sheet]
        rows = []
        for idx, rec in df.iterrows():
            for col, year in year_cols.items():
                volume = clean(rec[col])
                rows.append(
                    {
                        **provenance(ctx, doc, sheet=sheet, row=idx + 2, key=(col,)),
                        "sales_group": sheet,
                        "vehicle_codes": ",".join(config.SALES_SHEET_VEHICLES[sheet]),
                        "market": clean_str(rec["market"]),
                        "market_code": clean_str(rec["market_code"]),
                        "year": year,
                        "volume": int(volume) if volume is not None else None,
                        "is_planning_figure": year in config.SALES_PLANNING_YEARS,
                        "is_combined_models": len(vehicles) > 1,
                        "scope_label": (
                            f"{' + '.join(vehicles)} combined (not {vehicles[0]}-only)" if len(vehicles) > 1 else vehicles[0]
                        ),
                        "source_column": get_column_letter(list(df.columns).index(col) + 1),
                    }
                )
        result[sheet] = bulk_insert(ctx.session, SalesVolume, rows)
    return result


# ---------------------------------------------------------------------------
# PDF / PPTX text extraction
# ---------------------------------------------------------------------------


def ingest_pdf_pages(ctx: IngestionContext, doc: SourceDocument, pdf: pymupdf.Document) -> dict:
    rows, empty_pages = [], []
    for index, page in enumerate(pdf):
        text = page.get_text()
        images = len(page.get_images())
        note = None
        if not text.strip():
            empty_pages.append(index + 1)
            note = (
                f"No text layer on this page ({images} embedded images). "
                "Content is image-only; OCR was not performed, so no text is stored."
            )
        rows.append(
            {
                **provenance(ctx, doc, page=index + 1),
                "page_number": index + 1,
                "page_kind": "pdf_page",
                "text": text,
                "char_count": len(text.strip()),
                "image_count": images,
                "extraction_method": f"pymupdf {pymupdf.VersionBind} get_text",
                "extraction_note": note,
            }
        )
    return {"pages": bulk_insert(ctx.session, SourceDocumentPage, rows), "pages_without_text": empty_pages}


def ingest_pptx(ctx: IngestionContext, doc: SourceDocument, path: Path) -> dict:
    deck = Presentation(str(path))
    rows = []
    for number, slide in enumerate(deck.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        pictures = sum(1 for s in slide.shapes if s.shape_type == 13)
        text = "\n".join(parts)
        rows.append(
            {
                **provenance(ctx, doc, page=number),
                "page_number": number,
                "page_kind": "pptx_slide",
                "text": text,
                "char_count": len(text.strip()),
                "image_count": pictures,
                "extraction_method": "python-pptx text frames + tables",
                "extraction_note": (
                    "Text inside pictures, grouped shapes and charts is not extracted." if pictures else None
                ),
            }
        )
    return {"slides": bulk_insert(ctx.session, SourceDocumentPage, rows)}


# ---------------------------------------------------------------------------
# Option-list (equipment) PDF parsing
# ---------------------------------------------------------------------------

CODE_RE = re.compile(r"^[0-9A-Z]{2,4}$")
GROSS_RE = re.compile(r"^(?:\d{1,3}(?:\.\d{3})*,–|–,–)$")
NET_RE = re.compile(r"^\[\d{1,3}(?:\.\d{3})*,\d{2}\]$")
INLINE_PRICE_RE = re.compile(r"^(\d{1,3}(?:\.\d{3})*,–)\s*(\[\d{1,3}(?:\.\d{3})*,\d{2}\])?$")
CONDITION_RE = re.compile(r"^(Nicht in Verbindung|Nur in Verbindung|In Verbindung|Für )")
SYMBOL_AVAILABILITY = {"■": "standard", "□": "optional", "–,–": "no_extra_charge"}
FOOTER_Y = 780.0  # legend line ("■ Serienausstattung ...") sits below this
LINE_TOLERANCE = 1.6


@dataclass
class Span:
    x0: float
    y0: float
    x1: float
    y1: float
    size: float
    text: str

    @property
    def xc(self) -> float:
        return (self.x0 + self.x1) / 2


def page_spans(page: pymupdf.Page) -> list[Span]:
    spans = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            for s in line["spans"]:
                text = s["text"].replace("\xad", "").strip()
                if text:
                    spans.append(Span(*s["bbox"], s["size"], text))
    return spans


def parse_price(text: str | None) -> float | None:
    if not text or text.startswith("–"):
        return None
    return float(text.replace(".", "").replace(",–", "").replace(",", "."))


def cell_availability(raw: str) -> str:
    if raw in SYMBOL_AVAILABILITY:
        return SYMBOL_AVAILABILITY[raw]
    if GROSS_RE.match(raw.split(" ")[0]):
        return "optional_price_in_cell"
    return "other"


@dataclass
class ValueLine:
    y: float
    price_gross: str | None = None
    price_net: str | None = None
    cells: dict[str, str] = field(default_factory=dict)


def parse_option_page(page: pymupdf.Page) -> list[dict]:
    """Extract SA-coded equipment rows from one option-list page.

    Returns dicts with equipment fields plus an ``availability`` list. Pages
    without an 'SA-Nr.' column header yield nothing.
    """
    spans = [s for s in page_spans(page) if s.y0 < FOOTER_Y]
    sa_headers = [s for s in spans if s.text == "SA-Nr."]
    if not sa_headers:
        return []
    sa_x0 = min(s.x0 for s in sa_headers) - 8
    sa_x1 = max(s.x1 for s in sa_headers) + 8
    header_bottom = min(s.y1 for s in sa_headers) + 12

    title = next((s.text for s in spans if s.size >= 14), None)
    header_spans = [s for s in spans if s.y0 < header_bottom and s.x0 > sa_x1 - 8]
    price_header = [s for s in header_spans if "EUR" in s.text or "MwSt" in s.text or "%" in s.text]
    if not price_header:
        raise IngestionError(f"page {page.number + 1}: SA-Nr. column without price header")
    price_x0 = min(s.x0 for s in price_header) - 6
    price_x1 = max(s.x1 for s in price_header) + 6
    variants = sorted(
        (s for s in header_spans if s.x0 > price_x1 and s.size < 7 and s not in price_header),
        key=lambda s: s.xc,
    )

    def variant_for(span: Span) -> str | None:
        best = min(variants, key=lambda v: abs(v.xc - span.xc), default=None)
        return best.text if best and abs(best.xc - span.xc) <= 9 else None

    body = [s for s in spans if s.y0 >= header_bottom - 4 or s.x0 < sa_x0]
    codes: list[Span] = []
    for s in sorted(
        (s for s in body if s.x0 >= sa_x0 and s.x1 <= sa_x1 and CODE_RE.match(s.text) and s.y0 > header_bottom - 20),
        key=lambda s: (s.y0, s.x0),
    ):
        # Some codes are drawn twice (bold overlay); keep one.
        if not any(c.text == s.text and abs(c.y0 - s.y0) < 1 and abs(c.x0 - s.x0) < 2 for c in codes):
            codes.append(s)
    left = [s for s in body if s.x1 <= sa_x0 + 2 and s.text != "SA-Nr."]
    # Headings are upper-case ("FAHRASSISTENZ UND SICHERHEIT", "KUNSTLEDER").
    sections = [s for s in left if 7.5 <= s.size < 14 and s.text == s.text.upper()]
    subsections = [s for s in left if 6.5 <= s.size < 7.5 and s.text == s.text.upper()]

    # Value lines: prices and variant cells grouped by y; net prices "[...]"
    # printed on the following line are merged into the line above.
    lines: list[ValueLine] = []
    value_spans = sorted((s for s in body if s.x0 >= price_x0 and s.y0 >= header_bottom - 4), key=lambda s: (s.y0, s.x0))
    for s in value_spans:
        in_price_col = s.x1 <= price_x1 or not variants
        variant = None if in_price_col else variant_for(s)
        if not in_price_col and variant is None:
            continue
        line = lines[-1] if lines and abs(lines[-1].y - s.y0) <= LINE_TOLERANCE else None
        if line is None:
            line = ValueLine(y=s.y0)
            lines.append(line)
        if in_price_col:
            inline = INLINE_PRICE_RE.match(s.text)
            if inline:
                line.price_gross, line.price_net = inline.group(1), inline.group(2)
            elif GROSS_RE.match(s.text):
                line.price_gross = s.text
            elif NET_RE.match(s.text):
                line.price_net = s.text
        else:
            line.cells[variant] = s.text
    merged: list[ValueLine] = []
    for line in lines:
        only_net = (line.price_gross is None) and all(NET_RE.match(v) for v in line.cells.values())
        if merged and only_net and line.y - merged[-1].y <= 9:
            prev = merged[-1]
            prev.price_net = prev.price_net or line.price_net
            for variant, value in line.cells.items():
                if variant in prev.cells:
                    prev.cells[variant] = f"{prev.cells[variant]} {value}"
            continue
        if line.price_gross or line.price_net or line.cells:
            merged.append(line)

    def last_before(candidates: list[Span], y: float, floor: float = -1) -> str | None:
        found = [s for s in candidates if floor < s.y0 <= y + 0.5]
        return max(found, key=lambda s: s.y0).text if found else None

    def availability_rows(band_lines: list[ValueLine], band_left: list[Span], y0: float) -> list[dict]:
        rows = []
        for ln in band_lines:
            condition = None
            if abs(ln.y - y0) > LINE_TOLERANCE:
                cond = sorted((s for s in band_left if abs(s.y0 - ln.y) <= LINE_TOLERANCE), key=lambda s: s.x0)
                condition = " ".join(s.text for s in cond) or None
            for variant, raw in ln.cells.items():
                rows.append(
                    {
                        "variant": variant,
                        "condition_text": condition,
                        "line_price_gross_text": ln.price_gross,
                        "cell_raw": raw,
                        "availability": cell_availability(raw),
                        "line_y": round(ln.y, 1),
                    }
                )
        return rows

    headings = sections + subsections
    text_left = [s for s in left if s not in headings and s.size < 7.5]

    # A code may be printed up to ~5pt below or ~3pt above its name line (centred
    # across a two-line entry). Anchor each item at the highest left-column
    # line in that window; fall back to an upper-case heading on that line
    # (package names such as "PREMIUM PAKET" are set as headings).
    def anchor(code: Span) -> tuple[float, list[Span]]:
        window = [s for s in left if code.y0 - 5 <= s.y0 <= code.y0 + 4]
        pool = [s for s in window if s in text_left] or [s for s in window if s in headings]
        if not pool:
            return code.y0, []
        top = min(s.y0 for s in pool)
        return top, sorted((s for s in pool if abs(s.y0 - top) <= LINE_TOLERANCE), key=lambda s: s.x0)

    anchors = [anchor(c) for c in codes]

    results, group_code, group_index = [], None, None
    for i, code in enumerate(codes):
        anchor_y, name_spans = anchors[i]
        y0 = min(anchor_y, code.y0)
        heading_y = max(anchor_y, code.y0)
        next_y = anchors[i + 1][0] if i + 1 < len(codes) else FOOTER_Y
        next_heading = min((s.y0 for s in headings if s.y0 > y0 + LINE_TOLERANCE), default=FOOTER_Y)
        y_end = min(next_y, next_heading) - 0.5
        band_left = [s for s in text_left if y0 - LINE_TOLERANCE <= s.y0 < y_end]
        name_note = None
        if name_spans and CONDITION_RE.match(name_spans[0].text):
            # Only a condition line is printed level with the code; the option
            # name itself is not in the PDF text layer. Do not guess it.
            name_spans = []
            name_note = "Option name not present in the PDF text layer; see page_title."
        # Names wrapped onto a second line ("... Bicolor mit" / "Mischbereifung").
        if name_spans and re.search(r"(\bmit|\bund|-)$", name_spans[-1].text):
            following = [s for s in band_left if 0 < s.y0 - y0 <= 8 + LINE_TOLERANCE]
            name_spans += sorted(following, key=lambda s: s.x0)
        desc_spans = [s for s in band_left if s not in name_spans]
        band_lines = [ln for ln in merged if y0 - LINE_TOLERANCE <= ln.y < y_end]

        section_span_y = max((s.y0 for s in sections if s.y0 <= heading_y + 0.5), default=-1)
        if group_code and not code.text.startswith(group_code):
            group_code, group_index = None, None
        parent = group_code if group_code and code.text != group_code else None
        if not band_lines and len(code.text) == 2 and i + 1 < len(codes) and codes[i + 1].text.startswith(code.text):
            group_code, group_index = code.text, len(results)

        if parent is not None:
            # A group child owns only the line printed level with its code; later
            # lines (e.g. "In Verbindung mit M Sportpaket") apply to the group.
            own = [ln for ln in band_lines if abs(ln.y - code.y0) <= LINE_TOLERANCE]
            group_lines = [ln for ln in band_lines if ln not in own]
            if group_lines:
                parent_item = results[group_index]
                parent_item["availability"] += availability_rows(group_lines, band_left, y0=-100)
                parent_item["parse_note"] = (
                    "Availability lines printed after the colour group are attached to the group code."
                )
            band_lines = own

        first = band_lines[0] if band_lines else None
        availability = availability_rows(band_lines, band_left, code.y0)

        note = name_note
        if variants and not availability and code.text != group_code:
            note = (
                "No availability symbols on this code's line; the catalogue may show shared "
                "availability on a neighbouring line (e.g. colour group). Not inferred."
            )
        raw_text = " ".join(s.text for s in sorted(band_left, key=lambda s: (s.y0, s.x0)))
        results.append(
            {
                "page_title": title,
                "section": last_before(sections, heading_y),
                "subsection": last_before(subsections, heading_y, floor=section_span_y),
                "option_code": code.text,
                "parent_option_code": parent,
                "name": " ".join(s.text for s in name_spans) or None,
                "description": " ".join(s.text for s in sorted(desc_spans, key=lambda s: (s.y0, s.x0))) or None,
                "price_gross_eur_text": first.price_gross if first else None,
                "price_net_eur_text": first.price_net if first else None,
                "price_gross_eur": parse_price(first.price_gross) if first else None,
                "has_variant_matrix": bool(variants),
                "raw_text": raw_text or code.text,
                "parse_note": note,
                "y": round(y0, 1),
                "availability": availability,
            }
        )
    return results


def ingest_option_list(ctx: IngestionContext, doc: SourceDocument, path: Path, vehicle_code: str) -> dict:
    with pymupdf.open(path) as pdf:
        page_result = ingest_pdf_pages(ctx, doc, pdf)
        equipment_rows, availability_rows = [], []
        pages_with_equipment = 0
        for index, page in enumerate(pdf):
            items = parse_option_page(page)
            pages_with_equipment += bool(items)
            for item in items:
                avail = item.pop("availability")
                y = item.pop("y")
                prov = provenance(ctx, doc, page=index + 1, key=("equipment", item["option_code"], y))
                equipment_rows.append({**prov, "vehicle_code": vehicle_code, **item})
                for a in avail:
                    line_y = a.pop("line_y")
                    availability_rows.append(
                        {
                            **provenance(
                                ctx, doc, page=index + 1, key=("availability", item["option_code"], y, line_y, a["variant"])
                            ),
                            "equipment_record_id": prov["record_id"],
                            "vehicle_code": vehicle_code,
                            "option_code": item["option_code"],
                            **a,
                        }
                    )
    return {
        **page_result,
        "pages_with_equipment": pages_with_equipment,
        "equipment": bulk_insert(ctx.session, Equipment, equipment_rows),
        "equipment_availability": bulk_insert(ctx.session, EquipmentAvailability, availability_rows),
    }


# ---------------------------------------------------------------------------
# Consolidated workbook: overlap detection (not re-imported)
# ---------------------------------------------------------------------------

# consolidated sheet -> (original file, original sheet, header row or None)
CONSOLIDATED_SHEET_MAP: dict[str, tuple[str, str, int | None]] = {
    "G60 Feedback All": (config.FEEDBACK_FILES["G60"], config.FEEDBACK_SHEET, 0),
    "Public Competitors": (config.COMPETITOR_EVIDENCE_FILE, "Public Evidence", 0),
    "Public Caveats": (config.COMPETITOR_EVIDENCE_FILE, "Method and caveats", 0),
    "Sales G60 G68": (config.SALES_FILE, "G60_G68", 0),
    "Study CN EU 2025": (config.CUSTOMER_STUDIES_FILE, "CN_EU_2025", None),
    "Study US 2025": (config.CUSTOMER_STUDIES_FILE, "US_2025", None),
}


def _comparable(df: pd.DataFrame) -> pd.DataFrame:
    return df.map(lambda v: "" if clean(v) is None else str(v)).reset_index(drop=True)


def compare_frames(a: pd.DataFrame, b: pd.DataFrame) -> int | None:
    """Number of differing cells, or None if shapes/headers differ."""
    if a.shape != b.shape or [str(c) for c in a.columns] != [str(c) for c in b.columns]:
        return None
    a_, b_ = _comparable(a), _comparable(b)
    b_.columns = a_.columns
    return int((a_ != b_).to_numpy().sum())


def theme_counts_from_feedback(feedback: pd.DataFrame) -> pd.DataFrame:
    """Recompute the consolidated 'Feedback Theme Counts' sheet from the original
    G60 feedback rows (its labels for missing values are 'Unclassified' and
    'Unspecified'), so the derived sheet can be verified rather than trusted."""
    grouped = (
        feedback.assign(
            l2=feedback["Vfc level2 Name"].fillna("Unclassified"),
            ft=feedback["Feedback Type"].fillna("Unspecified"),
        )
        .groupby(["l2", "ft"])
        .size()
        .rename("Records")
        .reset_index()
    )
    return grouped.rename(columns={"l2": "VFC Level 2", "ft": "Feedback type"})


def analyse_consolidated_workbook(ctx: IngestionContext, doc: SourceDocument, path: Path) -> dict:
    sheets = excel_sheet_names(path)
    overlaps, summary = [], {}
    for sheet in sheets:
        cons_rows = None
        record = {
            "consolidated_file": path.name,
            "consolidated_sheet": sheet,
            "original_file": None,
            "original_sheet": None,
            "original_rows": None,
            "differing_cells": None,
            "checked_at": ctx.imported_at,
        }
        if sheet in CONSOLIDATED_SHEET_MAP:
            orig_file, orig_sheet, header = CONSOLIDATED_SHEET_MAP[sheet]
            orig_path = find_data_file(ctx, orig_file)
            cons = read_sheet(path, sheet, header=header)
            orig = read_sheet(orig_path, orig_sheet, header=header)
            cons_rows = len(cons)
            diff = compare_frames(cons, orig)
            record.update(original_file=orig_file, original_sheet=orig_sheet, original_rows=len(orig), differing_cells=diff)
            if diff == 0:
                result, decision = "exact_copy", (
                    f"Exact copy of {orig_file}/{orig_sheet}. Original is the source of truth; not re-imported."
                )
            else:
                result, decision = "differs", (
                    f"Differs from {orig_file}/{orig_sheet} ({'shape/header mismatch' if diff is None else f'{diff} cells'}). "
                    "Original kept as source of truth; consolidated copy NOT imported - review manually."
                )
        elif sheet == "Feedback Theme Counts":
            cons = read_sheet(path, sheet)
            cons_rows = len(cons)
            orig = read_sheet(find_data_file(ctx, config.FEEDBACK_FILES["G60"]), config.FEEDBACK_SHEET)
            recomputed = theme_counts_from_feedback(orig)
            key = ["VFC Level 2", "Feedback type"]
            joined = cons.merge(recomputed, on=key, how="outer", suffixes=("_cons", "_orig"), indicator=True)
            mismatch = int(
                ((joined["_merge"] != "both") | (joined["Records_cons"] != joined["Records_orig"])).sum()
            )
            record.update(
                original_file=config.FEEDBACK_FILES["G60"],
                original_sheet=config.FEEDBACK_SHEET,
                original_rows=len(recomputed),
                differing_cells=mismatch,
            )
            result = "derived_aggregate"
            decision = (
                "Aggregate of G60 feedback rows by VFC level 2 x feedback type. "
                + ("Recomputed from original and matches exactly. " if mismatch == 0 else f"{mismatch} group(s) do NOT match the recomputation. ")
                + "Not imported; aggregates are recomputed from feedback_classifications."
            )
        elif sheet == "Evidence Link Template":
            cons_rows = len(read_sheet(path, sheet))
            result, decision = "template", "Template with an explicit example row ('Example only'). Not data; not imported."
        elif sheet == "START HERE":
            cons_rows = ingest_note_sheet(ctx, doc, path, sheet)
            result, decision = "metadata", "Workbook description; stored in dataset_notes."
        else:
            cons_rows = len(read_sheet(path, sheet))
            result, decision = "no_original", "No matching original sheet identified. Not imported - review manually."
        overlaps.append({**record, "consolidated_rows": cons_rows, "comparison_result": result, "decision": decision})
        summary[sheet] = result
    bulk_insert(ctx.session, DatasetOverlap, overlaps)
    return summary


# ---------------------------------------------------------------------------
# Cross-dataset relationships among public evidence
# ---------------------------------------------------------------------------


def flag_shared_public_sources(ctx: IngestionContext) -> int:
    """Public feedback and competitor evidence rows that cite the same URL.
    These describe the same underlying source; flagged, never merged."""
    session = ctx.session
    by_url: dict[str, list[tuple[str, str, str]]] = {}
    for rid, pid, url in session.query(PublicFeedback.record_id, PublicFeedback.original_public_id, PublicFeedback.source_url):
        if url:
            by_url.setdefault(url.strip(), []).append(("public_feedback", rid, pid))
    for rid, row, url in session.query(CompetitorEvidence.record_id, CompetitorEvidence.source_row, CompetitorEvidence.source_url):
        if url:
            by_url.setdefault(url.strip(), []).append(("competitor_evidence", rid, f"row{row}"))
    rows = []
    for url, members in by_url.items():
        tables = {m[0] for m in members}
        if len(members) < 2:
            continue
        rows.append(
            {
                "table_name": "+".join(sorted(tables)),
                "rule": "same_source_url_across_public_datasets" if len(tables) > 1 else "same_source_url",
                "group_key": url,
                "record_ids": ",".join(m[1] for m in members),
                "original_ids": ",".join(f"{m[0]}:{m[2]}" for m in members),
                "record_count": len(members),
                "review_status": "unreviewed",
            }
        )
    return bulk_insert(session, PossibleDuplicate, rows)


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------


def find_data_file(ctx: IngestionContext, file_name: str) -> Path:
    assert ctx.data_dir is not None
    matches = list(ctx.data_dir.rglob(file_name))
    if not matches:
        raise FileNotFoundError(f"{file_name} not found under {ctx.data_dir}")
    if len(matches) > 1:
        raise IngestionError(f"{file_name} found more than once: {matches}")
    return matches[0]


def _run_step(ctx: IngestionContext, report: dict, file_name: str, origin: str, count_fn, step, *, sot: bool = True) -> None:
    """Register a document and run its ingestion step inside a savepoint; on
    failure, record the problem on the document instead of inventing data."""
    try:
        path = find_data_file(ctx, file_name)
    except (FileNotFoundError, IngestionError) as exc:
        report["files"][file_name] = {"status": "missing", "error": str(exc)}
        log.error("%s", exc)
        return
    doc = register_document(ctx, path, origin, sot, count_fn(path))
    savepoint = ctx.session.begin_nested()
    try:
        result = step(doc, path)
        savepoint.commit()
        if isinstance(result, dict) and result.get("pages_without_text"):
            missing = result["pages_without_text"]
            doc.ingestion_status = "no_text_layer" if len(missing) == result["pages"] else "partially_extracted"
            doc.ingestion_notes = (
                f"Pages without a text layer: {missing}. Image-only content; OCR not performed, no text invented."
            )
        report["files"][file_name] = {"status": doc.ingestion_status, "result": result}
    except Exception as exc:  # noqa: BLE001 - every failure is documented, not hidden
        savepoint.rollback()
        doc.ingestion_status = "failed"
        doc.ingestion_notes = f"{type(exc).__name__}: {exc}"
        report["files"][file_name] = {"status": "failed", "error": doc.ingestion_notes}
        log.exception("Ingestion failed for %s", file_name)


def _pdf_pages(path: Path) -> int:
    with pymupdf.open(path) as pdf:
        return pdf.page_count


def _sheet_count(path: Path) -> int:
    return len(excel_sheet_names(path))


def run_full_ingestion(session: Session, data_dir: Path) -> dict:
    from services.source_discovery import load_manifest

    ctx = IngestionContext(session=session, data_dir=data_dir, manifest=load_manifest(data_dir))
    report: dict = {"imported_at": ctx.imported_at.isoformat(), "files": {}}

    for code, name in config.FEEDBACK_FILES.items():
        _run_step(ctx, report, name, "internal_bmw", _sheet_count,
                  lambda doc, path, code=code: ingest_feedback_file(ctx, doc, path, code))
    _run_step(ctx, report, config.CUSTOMER_STUDIES_FILE, "internal_bmw", _sheet_count,
              lambda doc, path: ingest_customer_studies(ctx, doc, path))
    _run_step(ctx, report, config.SALES_FILE, "internal_bmw", _sheet_count,
              lambda doc, path: ingest_sales(ctx, doc, path))
    for code, name in config.OPTION_LIST_FILES.items():
        _run_step(ctx, report, name, "internal_bmw", _pdf_pages,
                  lambda doc, path, code=code: ingest_option_list(ctx, doc, path, code))
    _run_step(ctx, report, config.PUBLIC_FEEDBACK_FILE, "public_external", _sheet_count,
              lambda doc, path: ingest_public_feedback(ctx, doc, path))
    _run_step(ctx, report, config.COMPETITOR_EVIDENCE_FILE, "public_external", _sheet_count,
              lambda doc, path: ingest_competitor_evidence(ctx, doc, path))

    def _briefing_pdf(doc, path):
        with pymupdf.open(path) as pdf:
            return ingest_pdf_pages(ctx, doc, pdf)

    _run_step(ctx, report, config.BRIEFING_PDF_FILE, "briefing", _pdf_pages, _briefing_pdf)
    _run_step(ctx, report, config.KEYNOTE_PPTX_FILE, "briefing",
              lambda p: len(Presentation(str(p)).slides),
              lambda doc, path: ingest_pptx(ctx, doc, path))
    _run_step(ctx, report, config.CONSOLIDATED_FILE, "internal_bmw", _sheet_count,
              lambda doc, path: analyse_consolidated_workbook(ctx, doc, path), sot=False)
    consolidated = session.get(SourceDocument, make_record_id("document", config.CONSOLIDATED_FILE))
    if consolidated is not None and consolidated.ingestion_status != "failed":
        consolidated.ingestion_status = "registered_only"
        consolidated.ingestion_notes = (
            "Consolidated workbook. Overlapping sheets compared against originals (see dataset_overlaps); "
            "only its START HERE notes are stored. Originals are the single source of truth."
        )
        report["files"][config.CONSOLIDATED_FILE]["status"] = consolidated.ingestion_status

    report["public_shared_source_groups"] = flag_shared_public_sources(ctx)

    known = {d.file_name for d in session.query(SourceDocument)}
    unregistered = sorted(
        p.name for p in data_dir.rglob("*")
        if p.is_file() and not p.name.startswith(".") and p.name not in known and p.name != "source_manifest.json"
    )
    report["unregistered_files"] = unregistered
    session.flush()
    return report
