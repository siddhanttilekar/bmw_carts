"""Management reports (Part 6): Excel, CSV and PDF from saved, validated data.

* Built only from stored records of one prioritization run (+ optional saved
  scenarios, challenges, debates). Missing values stay blank; incomplete
  requirements are never ranked.
* Verbatim customer feedback is EXCLUDED by default (IDs only); it can be
  included explicitly. No credentials, environment values or database files
  are ever included.
* Every export carries the prototype-methodology disclaimer and is logged in
  management_report_exports (metadata only).
"""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import uuid
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from core.decision_models import ManagementReportExport
from core.models import Feedback
from core.prioritization_models import PriorityAssessment
from services.agent_debate import list_debates
from services.critical_issues import flags_for
from services.decision_grounding import bind_run, run_agents
from services.dynamic_prioritization import agent_results, ranking
from services.priority_challenges import list_challenges
from services.requirement_review import get_requirement, history
from services.what_if_simulator import load_scenario

TITLE = "BMW ProductMind AI - Requirement Prioritization Report"
SCOPE = "BMW 5 Series (G60) customer evidence -> next-generation BMW 5 Series requirements"
DISCLAIMER = ("Prototype: the scoring methodology, agents and weights are a hackathon prototype and NOT an official BMW "
              "product-management framework. Scores are computed from stored data; incomplete requirements are not ranked; "
              "a priority score does not mean a requirement is approved.")
CONFIDENTIALITY = "Contains internal BMW information. Verbatim customer feedback is excluded unless explicitly selected."


class ReportOptions(BaseModel):
    priority_statuses: list[str] = Field(default_factory=lambda: ["COMPLETE", "INCOMPLETE"])
    review_statuses: list[str] | None = None
    include_evidence_refs: bool = True
    include_verbatim_feedback: bool = False
    verbatim_limit_per_requirement: int = 5
    include_assessment_rationales: bool = True
    include_challenges: bool = False
    include_debates: bool = False
    scenario_ids: list[str] = Field(default_factory=list)
    include_review_history: bool = True


def _now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _v(x):
    return "" if x is None else x


def build_report_data(session: Session, run_id: str, options: ReportOptions | None = None) -> dict:
    options = options or ReportOptions()
    run = bind_run(session, run_id)
    snapshot = json.loads(run.profile_snapshot_json)
    rows = [r for r in ranking(session, run_id) if r.priority_status in options.priority_statuses]
    if options.review_statuses:
        rows = [r for r in rows if r.review_status in options.review_statuses]
    ids = [r.requirement_id for r in rows]

    meta = [("Report", TITLE), ("Scope", SCOPE), ("Generated (UTC)", f"{_now():%Y-%m-%d %H:%M}"),
            ("Profile", f"{snapshot['profile_name']} ({run.profile_id})"), ("Profile version", run.profile_version),
            ("Illustrative profile", "yes" if snapshot.get("illustrative") else "no"),
            ("Prioritization run", f"{run_id} created {run.created_at:%Y-%m-%d %H:%M} by {run.created_by}"),
            ("Agent configuration versions", "; ".join(f"{a['name']} v{a['configuration_version']}" for a in run_agents(run))),
            ("Scoring method versions", "; ".join(f"{k} v{v}" for k, v in json.loads(run.method_versions_json).items())),
            ("Requirements in report", len(ids)), ("Methodology", DISCLAIMER), ("Confidentiality", CONFIDENTIALITY)]

    requirements, agent_rows, assessments, assumptions, evidence, flags, review = [], [], [], [], [], [], []
    for r in rows:
        req = get_requirement(session, r.requirement_id)
        requirements.append({"requirement_id": r.requirement_id, "title": r.title, "current_version": req.current_version,
                             "review_status_at_run": r.review_status, "current_review_status": req.review_status,
                             "priority_status": r.priority_status,
                             "final_score": _v(r.final_score) if r.priority_status == "COMPLETE" else "",
                             "rank": _v(r.rank) if r.priority_status == "COMPLETE" else "",
                             "missing_agents": ", ".join(r.missing_agents),
                             "critical_flags": "; ".join(f"{f['category']} ({f['status']})" for f in flags_for(session, r.requirement_id)),
                             "source_cluster_id": req.source_cluster_id, "statement": req.proposed_requirement})
        for a in agent_results(session, run_id, r.requirement_id):
            refs = [x["key"] for x in a["source_refs"]]
            agent_rows.append({"requirement_id": r.requirement_id, "agent": a["agent_name"], "agent_version": a["agent_version"],
                               "scoring_method": a["scoring_method"], "method_version": a["scoring_version"],
                               "status": a["status"], "score": _v(a["score"]), "weight_percent": a["weight_percent"],
                               "weighted_contribution": _v(a["weighted_contribution"]),
                               "rationale": a["rationale"] if options.include_assessment_rationales or "assessment_id" not in a["inputs"]
                               else "(rationale excluded)",
                               "evidence_reference_count": len(refs),
                               "evidence_references": "; ".join(refs[:25]) + (f"; ... (+{len(refs) - 25})" if len(refs) > 25 else "")
                               if options.include_evidence_refs else "(excluded)"})
            aid = a["inputs"].get("assessment_id")
            if aid and (pa := session.get(PriorityAssessment, aid)):
                assessments.append({"requirement_id": r.requirement_id, "agent": a["agent_name"], "assessment_id": aid,
                                    "assessor_id": pa.assessor_id, "score": pa.score,
                                    "rationale": pa.rationale if options.include_assessment_rationales else "(excluded)",
                                    "assessed_at": f"{pa.assessment_timestamp:%Y-%m-%d %H:%M}",
                                    "superseded_since": "yes" if pa.superseded else "no"})
            for item in a["explanation"]["items"]:
                if item["kind"] in ("ASSUMPTION", "LIMITATION"):
                    assumptions.append({"requirement_id": r.requirement_id, "source": a["agent_name"], "kind": item["kind"],
                                        "text": item["text"]})
        for text in json.loads(req.assumptions_json):
            assumptions.append({"requirement_id": r.requirement_id, "source": "requirement", "kind": "ASSUMPTION", "text": text})
        if options.include_evidence_refs:
            from core.requirement_models import RequirementEvidenceLink

            links = session.scalars(select(RequirementEvidenceLink).where(
                RequirementEvidenceLink.requirement_id == r.requirement_id).order_by(RequirementEvidenceLink.link_id)).all()
            texts = {}
            if options.include_verbatim_feedback:
                direct = [l.source_key for l in links if l.tier == "DIRECT"][:options.verbatim_limit_per_requirement]
                texts = dict(session.execute(select(Feedback.original_feedback_id, Feedback.customer_feedback)
                                             .where(Feedback.original_feedback_id.in_(direct))).all())
            for l in links:
                evidence.append({"requirement_id": r.requirement_id, "tier": l.tier, "source_table": l.source_table,
                                 "source_key": l.source_key, "scope": l.evidence_scope,
                                 "requires_review": "yes" if l.requires_review else "no",
                                 "verbatim_text": texts.get(l.source_key, "") if options.include_verbatim_feedback else ""})
        for f in flags_for(session, r.requirement_id, include_dismissed=True):
            flags.append({"requirement_id": r.requirement_id, **{k: f[k] for k in ("flag_id", "category", "status", "rationale",
                                                                                  "raised_by", "decided_by", "decision_note")}})
        if options.include_review_history:
            for h in history(session, r.requirement_id):
                review.append({"requirement_id": r.requirement_id, "at": f"{h.created_at:%Y-%m-%d %H:%M}", "action": h.action,
                               "from": _v(h.previous_status), "to": h.new_status, "by": h.reviewer_id, "reason": _v(h.reason),
                               "version": h.version_number})
    weights = [{"agent": a["name"], "agent_id": a["agent_id"], "agent_version": a["configuration_version"],
                "execution_method": a["execution_method"], "scoring_method": a["scoring_method"],
                "weight_percent": a["weight_bp"] / 100} for a in run_agents(run)]
    scenarios = []
    for sid in options.scenario_ids:
        sim = load_scenario(session, sid)
        for row in sim.rows:
            if row.requirement_id in ids:
                scenarios.append({"scenario_id": sid, "requirement_id": row.requirement_id,
                                  "official_final": _v(row.original_final), "official_rank": _v(row.original_rank),
                                  "simulated_status": row.simulated_status, "simulated_final": _v(row.simulated_final),
                                  "simulated_rank": _v(row.simulated_rank), "label": "SIMULATION (not official)",
                                  "configuration": "; ".join(f"{c['name']} {c['weight_bp'] / 100}%" for c in sim.configuration)})
    challenges = []
    if options.include_challenges:
        for c in list_challenges(session, run_id=run_id):
            if c["requirement_id"] in ids:
                challenges.append({"challenge_id": c["challenge_id"], "requirement_id": c["requirement_id"],
                                   "type": c["challenge_type"], "status": c["status"], "argument": c["argument"],
                                   "raised_by": c["raised_by"], "needs_review": "; ".join(c["needs_review"]),
                                   "resolution": _v(c["resolution_note"])})
    debates = []
    if options.include_debates:
        for d in list_debates(session, run_id=run_id):
            if d["requirement_a"] in ids and d["requirement_b"] in ids:
                debates.append({"debate_id": d["debate_id"], "requirement_a": d["requirement_a"],
                                "requirement_b": d["requirement_b"], "agreement": d["synthesis"]["agreement"],
                                "ranking_under_profile": d["synthesis"]["ranking_under_profile"],
                                "sensitivity": d["synthesis"]["sensitivity"]["classification"]})
    tables = {"Requirements": requirements, "Agent scores": agent_rows, "Weights": weights, "Human assessments": assessments,
              "Evidence": evidence, "Assumptions & limitations": assumptions, "Critical flags": flags,
              "Scenarios": scenarios, "Challenges": challenges, "Debates": debates, "Review history": review}
    return {"run_id": run_id, "meta": meta, "tables": tables, "options": options.model_dump()}


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------


def to_excel(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    ws.append(["Field", "Value"])
    for k, v in data["meta"]:
        ws.append([k, str(v)])
    _style(ws)
    for name, rows in data["tables"].items():
        if not rows and name in ("Scenarios", "Challenges", "Debates"):
            continue
        sheet = wb.create_sheet(name[:31])
        headers = list(rows[0]) if rows else ["(no records)"]
        sheet.append(headers)
        for r in rows:
            sheet.append([r[h] if not isinstance(r[h], (list, dict)) else json.dumps(r[h]) for h in headers])
        _style(sheet)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _style(ws) -> None:
    head = PatternFill("solid", fgColor="1E2C4A")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(12, width + 2), 70)
        for c in col[1:]:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def to_csv(data: dict) -> bytes:
    """Long format: one row per requirement x agent, with requirement fields and
    source references; run/profile/methodology columns on every row."""
    meta = dict(data["meta"])
    reqs = {r["requirement_id"]: r for r in data["tables"]["Requirements"]}
    fields = ["run_id", "profile", "profile_version", "requirement_id", "title", "review_status_at_run", "priority_status",
              "final_score", "rank", "critical_flags", "agent", "agent_version", "scoring_method", "method_version", "status",
              "score", "weight_percent", "weighted_contribution", "evidence_reference_count", "evidence_references",
              "methodology_note"]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fields)
    w.writeheader()
    for a in data["tables"]["Agent scores"]:
        r = reqs[a["requirement_id"]]
        w.writerow({"run_id": data["run_id"], "profile": meta["Profile"], "profile_version": meta["Profile version"],
                    "requirement_id": r["requirement_id"], "title": r["title"], "review_status_at_run": r["review_status_at_run"],
                    "priority_status": r["priority_status"], "final_score": r["final_score"], "rank": r["rank"],
                    "critical_flags": r["critical_flags"], **{k: a[k] for k in ("agent", "agent_version", "scoring_method",
                                                                               "method_version", "status", "score",
                                                                               "weight_percent", "weighted_contribution",
                                                                               "evidence_reference_count", "evidence_references")},
                    "methodology_note": "Prototype methodology - not an official BMW framework"})
    return buf.getvalue().encode("utf-8-sig")


def to_pdf(data: dict) -> bytes:
    """Readable landscape A4 report (ReportLab): wrapped cells, header rows
    repeated across pages, page footer with run ID and disclaimer."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    base = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=base["Title"], fontSize=16, leading=20, alignment=0, textColor=colors.HexColor("#0B1426"))
    h2 = ParagraphStyle("h2", parent=base["Heading2"], fontSize=12, leading=15, textColor=colors.HexColor("#0B1426"),
                        spaceBefore=10, spaceAfter=4)
    h3 = ParagraphStyle("h3", parent=base["Heading3"], fontSize=9.5, leading=12, spaceBefore=6, spaceAfter=2)
    body = ParagraphStyle("b", parent=base["BodyText"], fontSize=8.5, leading=11)
    cell = ParagraphStyle("c", parent=body, fontSize=7.5, leading=9.5)
    head = ParagraphStyle("hc", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    note = ParagraphStyle("n", parent=body, backColor=colors.HexColor("#F2F5FA"), borderColor=colors.HexColor("#C9D3E3"),
                          borderWidth=0.5, borderPadding=5, spaceBefore=4, spaceAfter=8)
    page_w = landscape(A4)[0] - 30 * mm

    def esc(x) -> str:
        return html.escape("" if x is None else str(x))

    def table(rows, cols, labels, widths, limit=None):
        if not rows:
            return [Paragraph("No records.", body)]
        shown = rows[:limit] if limit else rows
        data_rows = [[Paragraph(esc(l), head) for l in labels]]
        data_rows += [[Paragraph(esc(r.get(c, "")), cell) for c in cols] for r in shown]
        t = LongTable(data_rows, colWidths=[page_w * w for w in widths], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E2C4A")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#DDDDDD")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FC")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        out = [t]
        if limit and len(rows) > limit:
            out.append(Paragraph(f"... {len(rows) - limit} more rows in the Excel export.", body))
        return out

    t, meta = data["tables"], data["meta"]
    ranked = [r for r in t["Requirements"] if r["priority_status"] == "COMPLETE"]
    incomplete = [r for r in t["Requirements"] if r["priority_status"] != "COMPLETE"]
    story = [Paragraph(esc(TITLE), h1), Paragraph(esc(SCOPE), body), Spacer(1, 4),
             Paragraph(f"<b>Methodology.</b> {esc(DISCLAIMER)}<br/><b>Confidentiality.</b> {esc(CONFIDENTIALITY)}", note),
             Paragraph("Run and configuration", h2)]
    story += table([{"k": k, "v": v} for k, v in meta if k not in ("Report", "Methodology", "Confidentiality")],
                   ["k", "v"], ["Field", "Value"], [0.22, 0.78])
    story += [Paragraph("Agent weights", h2)]
    story += table(t["Weights"], ["agent", "agent_version", "execution_method", "scoring_method", "weight_percent"],
                   ["Agent", "Version", "Execution", "Scoring method", "Weight %"], [0.3, 0.08, 0.2, 0.3, 0.12])
    story += [Paragraph(f"Ranked requirements ({len(ranked)})", h2)]
    story += table(ranked, ["rank", "requirement_id", "title", "final_score", "current_review_status", "critical_flags"],
                   ["Rank", "Requirement", "Title", "Final score", "Review status", "Critical flags"],
                   [0.06, 0.15, 0.37, 0.1, 0.12, 0.2])
    story += [Paragraph(f"Incomplete requirements - not ranked ({len(incomplete)})", h2)]
    story += table(incomplete, ["requirement_id", "title", "missing_agents", "current_review_status", "critical_flags"],
                   ["Requirement", "Title", "Missing agents", "Review status", "Critical flags"],
                   [0.15, 0.3, 0.3, 0.1, 0.15])
    confirmed = [f for f in t["Critical flags"] if f["status"] == "CONFIRMED"]
    if confirmed:
        story += [Paragraph("Confirmed critical issues (shown regardless of rank)", h2)]
        story += table(confirmed, ["requirement_id", "category", "rationale", "decided_by"],
                       ["Requirement", "Category", "Rationale", "Confirmed by"], [0.15, 0.15, 0.55, 0.15])
    story += [Paragraph("Agent score breakdown", h2)]
    for r in t["Requirements"]:
        rows = [a for a in t["Agent scores"] if a["requirement_id"] == r["requirement_id"]]
        score = f", score {r['final_score']}, rank {r['rank']}" if r["final_score"] != "" else ""
        story += [Paragraph(f"{esc(r['requirement_id'])} - {esc(r['title'])} ({esc(r['priority_status'])}{score})", h3)]
        story += table(rows, ["agent", "status", "score", "weight_percent", "weighted_contribution",
                              "evidence_reference_count", "rationale"],
                       ["Agent", "Status", "Score", "Weight %", "Contribution", "Refs", "Rationale / calculation"],
                       [0.18, 0.09, 0.07, 0.07, 0.09, 0.05, 0.45])
    if t["Human assessments"]:
        story += [Paragraph("Human assessments", h2)]
        story += table(t["Human assessments"], ["requirement_id", "agent", "assessor_id", "score", "assessed_at", "rationale"],
                       ["Requirement", "Agent", "Assessor", "Score", "When", "Rationale"], [0.14, 0.18, 0.1, 0.06, 0.12, 0.4])
    if t["Scenarios"]:
        story += [Paragraph("Saved scenarios (SIMULATION - not official)", h2)]
        story += table(t["Scenarios"], ["scenario_id", "requirement_id", "official_final", "official_rank", "simulated_status",
                                        "simulated_final", "simulated_rank", "configuration"],
                       ["Scenario", "Requirement", "Official", "Rank", "Sim. status", "Simulated", "Sim. rank", "Configuration"],
                       [0.12, 0.14, 0.07, 0.05, 0.09, 0.08, 0.06, 0.39])
    if t["Challenges"]:
        story += [Paragraph("Priority challenges", h2)]
        story += table(t["Challenges"], ["challenge_id", "requirement_id", "type", "status", "argument", "needs_review", "resolution"],
                       ["Challenge", "Requirement", "Type", "Status", "Argument", "Needs review", "Resolution"],
                       [0.11, 0.12, 0.07, 0.09, 0.22, 0.25, 0.14])
    if t["Debates"]:
        story += [Paragraph("Agent debates", h2)]
        story += table(t["Debates"], ["requirement_a", "requirement_b", "agreement", "ranking_under_profile", "sensitivity"],
                       ["A", "B", "Agreement", "Ranking under profile", "Sensitivity"], [0.12, 0.12, 0.28, 0.3, 0.18])
    story += [Paragraph("Assumptions and limitations", h2)]
    story += table(t["Assumptions & limitations"], ["requirement_id", "kind", "source", "text"],
                   ["Requirement", "Kind", "Source", "Text"], [0.14, 0.1, 0.16, 0.6], limit=60)
    if t["Evidence"]:
        story += [Paragraph("Evidence references", h2),
                  Paragraph("Direct = cluster members (G60 feedback IDs); supporting/contextual are linked at "
                            "need-category level. Full list in the Excel export.", body)]
        story += table(t["Evidence"], ["requirement_id", "tier", "source_table", "source_key", "scope", "verbatim_text"],
                       ["Requirement", "Tier", "Table", "Reference", "Scope", "Verbatim (only if selected)"],
                       [0.13, 0.09, 0.12, 0.22, 0.1, 0.34], limit=80)
    if t["Review history"]:
        story += [Paragraph("Decision and approval history", h2)]
        story += table(t["Review history"], ["requirement_id", "at", "action", "from", "to", "by", "reason"],
                       ["Requirement", "When", "Action", "From", "To", "By", "Reason"],
                       [0.13, 0.11, 0.08, 0.11, 0.11, 0.1, 0.36])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(15 * mm, 8 * mm, f"BMW ProductMind AI · run {data['run_id']} · prototype methodology - not an "
                                           f"official BMW framework · page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm, topMargin=12 * mm,
                            bottomMargin=14 * mm, title=TITLE, subject=SCOPE, author="BMW ProductMind AI (prototype)")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def export_report(session: Session, run_id: str, fmt: str, created_by: str,
                  options: ReportOptions | None = None) -> tuple[bytes, str]:
    if fmt not in ("xlsx", "csv", "pdf"):
        raise ValueError("Format must be xlsx, csv or pdf")
    if not (created_by or "").strip():
        raise ValueError("A user identifier is required.")
    options = options or ReportOptions()
    data = build_report_data(session, run_id, options)
    content = {"xlsx": to_excel, "csv": to_csv, "pdf": to_pdf}[fmt](data)
    name = f"productmind_priority_report_{run_id}_{_now():%Y%m%d_%H%M}.{fmt}"
    session.add(ManagementReportExport(export_id=f"EXP-{uuid.uuid4().hex[:10].upper()}", report_format=fmt, run_id=run_id,
                                       options_json=options.model_dump_json(), file_name=name,
                                       sha256=hashlib.sha256(content).hexdigest(), size_bytes=len(content),
                                       created_by=created_by.strip(), created_at=_now()))
    session.flush()
    return content, name
